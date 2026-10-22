"""Tests for the Dashboard 4G PowerPoint export (native editable charts in FB Share template).

All charts in the exported deck are NATIVE PowerPoint Line Charts (XL_CHART_TYPE.LINE),
NOT images, built inside the fb_share.pptx presentation template.
"""
from __future__ import annotations

import io
import re

import pytest
from pptx import Presentation

from app import create_app
from app.services.dashboard_pptx import (
    PayloadError,
    build_dashboard_pptx,
    parse_hex_color,
)

PPTX_MIME = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
ENDPOINT = "/api/dashboard_4g/export-pptx"


def sample_chart(title="Payload Total (GB)", count=7):
    return {
        "title": title,
        "categories": [f"2026-10-{14+i:02d}" for i in range(count)],
        "series": [
            {"name": "Total", "values": [100.0 + i * 5 for i in range(count)], "color": "#2563eb"}
        ]
    }


def sample_payload(*, charts=None, per_slide=4, **overrides):
    body = {
        "dashboard": "4G",
        "section": "KPI Trend",
        "per_slide": per_slide,
        "scope": {
            "date_range": "2026-10-14 to 2026-10-20",
            "granularity": "DAILY",
            "filter": "Site ID · TIM012, TIM149"
        },
        "charts": charts or [sample_chart("Payload Total (GB)"), sample_chart("DL Throughput (Mbps)")]
    }
    body.update(overrides)
    return body


def open_deck(blob: bytes) -> Presentation:
    return Presentation(io.BytesIO(blob))


def slide_charts(slide):
    return [shape.chart for shape in slide.shapes if shape.has_chart]


def slide_pictures(slide):
    return [shape for shape in slide.shapes if shape.shape_type == 13]


def slide_text(slide) -> str:
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


# ── Fixtures ────────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def app():
    application = create_app()
    application.config["TESTING"] = True
    return application


def _csrf_token(client) -> str:
    html = client.get("/login").get_data(as_text=True)
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match, "csrf token missing from login page"
    return match.group(1)


@pytest.fixture()
def auth_client(app):
    client = app.test_client()
    token = _csrf_token(client)
    with client.session_transaction() as sess:
        sess["username"] = "test_engineer"
        sess["role"] = "admin"
    return client, token


# ── Service tests: Native charts & FB Share template ───────────────────────

def test_color_parsing():
    c1 = parse_hex_color("#2563eb")
    assert c1 is not None and (c1[0], c1[1], c1[2]) == (0x25, 0x63, 0xEB)
    c2 = parse_hex_color("rgb(220, 38, 38)")
    assert c2 is not None and (c2[0], c2[1], c2[2]) == (220, 38, 38)
    assert parse_hex_color(None) is None
    assert parse_hex_color("invalid") is None


def test_build_deck_has_native_charts_and_no_pictures():
    charts = [sample_chart(f"KPI {i}") for i in range(4)]
    blob, meta = build_dashboard_pptx(sample_payload(charts=charts, per_slide=4))

    assert blob.startswith(b"PK"), "pptx must be a zip container"
    assert meta == {"charts": 4, "slides": 1}

    deck = open_deck(blob)
    assert len(deck.slides) == 1
    slide = deck.slides[0]

    # Must have 4 native charts, and ZERO pictures
    charts_found = slide_charts(slide)
    pictures_found = slide_pictures(slide)
    assert len(charts_found) == 4, "Should have 4 native charts on the slide"
    assert len(pictures_found) == 0, "Should have 0 picture/image shapes"

    # Header title check - updated to Dashboard 4G | KPI Monitoring
    text = slide_text(slide)
    assert "Dashboard 4G" in text
    assert "KPI Monitoring |" in text
    assert "KPI Trend" in text
    assert "Network Benchmarking" not in text

    # Verify each chart has real categories and series data
    ch0 = charts_found[0]
    assert ch0.chart_title.text_frame.text == "KPI 0"
    assert len(list(ch0.plots[0].categories)) == 7
    assert len(ch0.plots[0].series) == 1
    assert ch0.plots[0].series[0].name == "Total"
    assert list(ch0.plots[0].series[0].values) == [100.0, 105.0, 110.0, 115.0, 120.0, 125.0, 130.0]


def test_multi_slide_chunking():
    # 7 charts with per_slide=4 -> 2 slides: 4 on slide 1, 3 on slide 2
    charts = [sample_chart(f"KPI {i}") for i in range(7)]
    blob, meta = build_dashboard_pptx(sample_payload(charts=charts, per_slide=4))

    assert meta == {"charts": 7, "slides": 2}
    deck = open_deck(blob)
    assert len(deck.slides) == 2
    assert len(slide_charts(deck.slides[0])) == 4
    assert len(slide_charts(deck.slides[1])) == 3
    assert len(slide_pictures(deck.slides[0])) == 0
    assert len(slide_pictures(deck.slides[1])) == 0


def test_nine_and_sixteen_charts_per_slide():
    # Test 9 charts per slide
    charts9 = [sample_chart(f"KPI {i}") for i in range(9)]
    blob9, meta9 = build_dashboard_pptx(sample_payload(charts=charts9, per_slide=9))
    assert meta9 == {"charts": 9, "slides": 1}
    deck9 = open_deck(blob9)
    assert len(slide_charts(deck9.slides[0])) == 9
    assert len(slide_pictures(deck9.slides[0])) == 0

    # Test 16 charts per slide
    charts16 = [sample_chart(f"KPI {i}") for i in range(16)]
    blob16, meta16 = build_dashboard_pptx(sample_payload(charts=charts16, per_slide=16))
    assert meta16 == {"charts": 16, "slides": 1}
    deck16 = open_deck(blob16)
    assert len(slide_charts(deck16.slides[0])) == 16
    assert len(slide_pictures(deck16.slides[0])) == 0


def test_multi_series_chart_support():
    comparison_chart = {
        "title": "Hourly Traffic Comparison (GB)",
        "categories": ["00:00", "01:00", "02:00"],
        "series": [
            {"name": "Before", "values": [10.5, 12.0, 15.2], "color": "#00e676"},
            {"name": "After", "values": [14.0, 15.5, 18.0], "color": "#4facfe"}
        ]
    }
    blob, meta = build_dashboard_pptx(sample_payload(charts=[comparison_chart], per_slide=1))
    deck = open_deck(blob)
    ch = slide_charts(deck.slides[0])[0]
    assert len(ch.plots[0].series) == 2
    assert ch.plots[0].series[0].name == "Before"
    assert ch.plots[0].series[1].name == "After"
    assert ch.has_legend is True

    # Style verification: smooth lines, exact dashboard colors, font, and axes
    s0 = ch.plots[0].series[0]
    s1 = ch.plots[0].series[1]
    assert s0.smooth is True
    assert s1.smooth is True
    assert s0.format.line.color.rgb == (0x00, 0xE6, 0x76)  # Before neon green
    assert s1.format.line.color.rgb == (0x4F, 0xAC, 0xFE)  # After electric cyan
    assert ch.chart_title.text_frame.paragraphs[0].font.name == "Poppins"
    assert ch.value_axis.has_major_gridlines is False  # Clean view without gridlines
    assert ch.value_axis.has_title is True
    assert ch.value_axis.axis_title.text_frame.text == "GB"


def test_dual_axis_chart_generation():
    chart_dual = {
        "title": "Throughput vs Payload",
        "categories": ["14 Oct", "15 Oct", "16 Oct"],
        "series": [
            {"name": "DL Thp (Mbps)", "values": [25.5, 30.2, 28.0], "color": "#4facfe", "axis": "L"},
            {"name": "Payload (GB)", "values": [1200.0, 1500.0, 1400.0], "color": "#00e676", "axis": "R"}
        ],
        "left_axis_title": "Mbps",
        "right_axis_title": "GB"
    }
    blob, meta = build_dashboard_pptx(sample_payload(charts=[chart_dual], per_slide=1))
    assert meta == {"charts": 1, "slides": 1}

    deck = open_deck(blob)
    chart = deck.slides[0].shapes[2].chart  # shape 0 is title, shape 1 is card, shape 2 is chart
    assert len(chart.plots) == 2, "Dual axis chart must generate 2 separate plots (L and R)"
    assert chart.plots[0].series[0].name == "DL Thp (Mbps)"
    assert chart.plots[1].series[0].name == "Payload (GB)"


def test_simplify_hourly_categories_to_daily():
    from app.services.dashboard_pptx import simplify_hourly_categories

    hourly_cats = [
        "2026-10-14 00:00", "2026-10-14 01:00", "2026-10-14 02:00",
        "2026-10-15 00:00", "2026-10-15 01:00"
    ]
    simplified = simplify_hourly_categories(hourly_cats)
    assert simplified == ["14 Oct", " ", " ", "15 Oct", " "]

    daily_cats = ["2026-10-14", "2026-10-15"]
    assert simplify_hourly_categories(daily_cats) == ["14 Oct", "15 Oct"]


def test_templates_endpoint(auth_client):
    client, _ = auth_client
    res = client.get("/api/dashboard_4g/export-pptx-templates")
    assert res.status_code == 200
    data = res.get_json()
    assert data["status"] == "success"
    templates = data["templates"]
    assert len(templates) >= 2
    assert any(t["id"] == "infra_2026" for t in templates)


def test_empty_or_invalid_charts_rejected():
    with pytest.raises(PayloadError, match="select at least one chart"):
        build_dashboard_pptx({"charts": []})
    with pytest.raises(PayloadError, match="select at least one chart"):
        build_dashboard_pptx({})


# ── Endpoint tests ──────────────────────────────────────────────────────────

def test_endpoint_returns_real_pptx_with_native_charts(auth_client):
    client, token = auth_client
    body = sample_payload()
    res = client.post(ENDPOINT, json=body, headers={"X-CSRFToken": token})

    assert res.status_code == 200, res.get_data(as_text=True)[:500]
    assert res.mimetype == PPTX_MIME
    assert res.headers["Content-Disposition"].startswith("attachment;")
    assert ".pptx" in res.headers["Content-Disposition"]
    assert res.headers["X-Pptx-Charts"] == "2"
    assert res.headers["X-Pptx-Slides"] == "1"

    deck = open_deck(res.data)
    assert len(deck.slides) == 1
    assert len(slide_charts(deck.slides[0])) == 2
    assert len(slide_pictures(deck.slides[0])) == 0


def test_endpoint_rejects_empty_charts_with_json(auth_client):
    client, token = auth_client
    res = client.post(ENDPOINT, json={"charts": []}, headers={"X-CSRFToken": token})
    assert res.status_code == 400
    data = res.get_json()
    assert data["status"] == "error"
    assert "select at least one chart" in data["message"].lower()


def test_endpoint_rejects_missing_csrf(app):
    client = app.test_client()
    _csrf_token(client)
    with client.session_transaction() as sess:
        sess["username"] = "test_engineer"
    res = client.post(ENDPOINT, json=sample_payload())
    assert res.status_code == 400
    assert b"CSRF" in res.data or b"csrf" in res.data.lower()


def test_endpoint_requires_login(app):
    client = app.test_client()
    token = _csrf_token(client)
    res = client.post(ENDPOINT, json=sample_payload(), headers={"X-CSRFToken": token})
    assert res.status_code == 302
    assert "/login" in res.headers["Location"]


def test_dashboard_page_exposes_dedicated_pptx_modal_and_button(auth_client):
    client, _ = auth_client
    html = client.get("/dashboard_4g").get_data(as_text=True)
    assert 'id="exportPptxModal"' in html
    assert 'openPptxModal' in html
    assert 'executePptxExport' in html
    assert "/api/dashboard_4g/export-pptx" in html
