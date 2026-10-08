from flask import Blueprint, render_template, request, session, jsonify, send_file, make_response
from app import csrf
from app.db.db_webapp import get_postgres_connection
from ._utils import login_required, viewer_blocked, db_query
import logging
import time
import io
import csv
import psycopg2.extras
import datetime

logger = logging.getLogger(__name__)

sow = Blueprint("sow", __name__)

# In-memory caches for fast responses
_LOCATIONS_CACHE = {}
_KAB_REGIONS_CACHE = {"data": None, "timestamp": 0}
_YEARWEEKS_CACHE = {"data": None, "timestamp": 0}
_SUMMARY_YEARS_CACHE = {"data": None, "timestamp": 0}
_CACHE_TTL = 1800  # 30 minutes


def clear_sow_caches():
    """Clear memory caches after data updates/imports."""
    global _LOCATIONS_CACHE, _KAB_REGIONS_CACHE, _YEARWEEKS_CACHE, _SUMMARY_YEARS_CACHE
    _LOCATIONS_CACHE.clear()
    _KAB_REGIONS_CACHE = {"data": None, "timestamp": 0}
    _YEARWEEKS_CACHE = {"data": None, "timestamp": 0}
    _SUMMARY_YEARS_CACHE = {"data": None, "timestamp": 0}
    _FB_YEARWEEKS_CACHE["data"] = None
    _FB_YEARWEEKS_CACHE["timestamp"] = 0
    _FB_CITY_MASTER_CACHE["data"] = None
    _FB_CITY_MASTER_CACHE["timestamp"] = 0


def get_cached_kab_regions():
    """Retrieve distinct Region values for level='Kabupaten'."""
    now = time.time()
    if _KAB_REGIONS_CACHE["data"] and (now - _KAB_REGIONS_CACHE["timestamp"] < _CACHE_TTL):
        return _KAB_REGIONS_CACHE["data"]
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(
                """
                SELECT DISTINCT "Region"
                FROM sow.onx_rank_score
                WHERE level = 'Kabupaten' AND "Region" IS NOT NULL AND "Region" != '' AND "Region" != '-'
                ORDER BY "Region" ASC
                """
            )
            regions = [r[0] for r in cur.fetchall() if r[0]]
            _KAB_REGIONS_CACHE["data"] = regions
            _KAB_REGIONS_CACHE["timestamp"] = now
            return regions
    except Exception as e:
        logger.exception("Error loading Kabupaten regions: %s", e)
        return []


def get_cached_locations(level, region=None):
    """Retrieve distinct locations for a level (and optional Region for Kabupaten)."""
    cache_key = f"{level}:{region or ''}"
    if cache_key in _LOCATIONS_CACHE:
        return _LOCATIONS_CACHE[cache_key]
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            if level == "Kabupaten" and region:
                cur.execute(
                    """
                    SELECT DISTINCT location 
                    FROM sow.onx_rank_score 
                    WHERE level = %s AND "Region" = %s
                    ORDER BY location ASC
                    """,
                    (level, region)
                )
            else:
                cur.execute(
                    """
                    SELECT DISTINCT location 
                    FROM sow.onx_rank_score 
                    WHERE level = %s 
                    ORDER BY location ASC
                    """,
                    (level,)
                )
            locs = [r[0] for r in cur.fetchall() if r[0] is not None]
            _LOCATIONS_CACHE[cache_key] = locs
            return locs
    except Exception as e:
        logger.exception("Error loading locations for level %s, region %s: %s", level, region, e)
        return []


def get_cached_yearweeks():
    """Retrieve all distinct yearweeks sorted descending."""
    now = time.time()
    if _YEARWEEKS_CACHE["data"] and (now - _YEARWEEKS_CACHE["timestamp"] < _CACHE_TTL):
        return _YEARWEEKS_CACHE["data"]
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(
                """
                SELECT DISTINCT yearweek 
                FROM sow.onx_rank_score 
                ORDER BY yearweek DESC
                """
            )
            yws = [int(r[0]) for r in cur.fetchall() if r[0] is not None]
            _YEARWEEKS_CACHE["data"] = yws
            _YEARWEEKS_CACHE["timestamp"] = now
            return yws
    except Exception as e:
        logger.exception("Error loading yearweeks: %s", e)
        return []


# Standard 16 metrics in order
ORDERED_16_METRICS = [
    # Col 1: Overall experience & timeon
    "videoexperience_overall",
    "gamesexperience_overall",
    "voiceexperience_overall",
    "timeon_overall",
    # Col 2: Overall speed & quality
    "download_overall",
    "upload_overall",
    "reliability_overall",
    "consistentquality_overall",
    # Col 3: 5G experience & timeon
    "videoexperience_5g",
    "gamesexperience_5g",
    "voiceexperience_5g",
    "timeon_5g",
    # Col 4: 5G speed & coverage
    "download_5g",
    "upload_5g",
    "onxcoveragesim_5g",
    "onxcoveragesim_overall"
]

METRIC_TITLES = {
    "consistentquality_overall": "Consistent Quality Overall",
    "download_5g": "Download Speed 5G",
    "download_overall": "Download Speed Overall",
    "gamesexperience_5g": "Games Experience 5G",
    "gamesexperience_overall": "Games Experience Overall",
    "onxcoveragesim_5g": "Coverage Experience 5G",
    "onxcoveragesim_overall": "Coverage Experience Overall",
    "reliability_overall": "Reliability Experience",
    "timeon_5g": "Time on 5G",
    "timeon_overall": "Time on Overall",
    "upload_5g": "Upload Speed 5G",
    "upload_overall": "Upload Speed Overall",
    "videoexperience_5g": "Video Experience 5G",
    "videoexperience_overall": "Video Experience Overall",
    "voiceexperience_5g": "Voice Experience 5G",
    "voiceexperience_overall": "Voice Experience Overall",
}


# ══════════════════════════════════════════════════════════════════════
# FB SHARE (sow.fb_share) helpers
# ══════════════════════════════════════════════════════════════════════
_FB_YEARWEEKS_CACHE = {"data": None, "timestamp": 0}
FB_PROVIDERS = ("tsel", "ioh", "xls")  # tsel=telkomsel, ioh=isat3, xls=xl+smartfren


def get_cached_fb_yearweeks():
    """Distinct yearweeks available in sow.fb_share (descending)."""
    now = time.time()
    if _FB_YEARWEEKS_CACHE["data"] and (now - _FB_YEARWEEKS_CACHE["timestamp"] < _CACHE_TTL):
        return _FB_YEARWEEKS_CACHE["data"]
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute("SELECT DISTINCT yearweek FROM sow.fb_share WHERE yearweek IS NOT NULL ORDER BY yearweek DESC")
            yws = [int(r[0]) for r in cur.fetchall()]
            _FB_YEARWEEKS_CACHE["data"] = yws
            _FB_YEARWEEKS_CACHE["timestamp"] = now
            return yws
    except Exception as e:
        logger.exception("Error loading fb_share yearweeks: %s", e)
        return []


def _shift_yearweek(yw, weeks):
    """Shift an ISO yearweek (YYYYWW) back by N weeks. Returns None if invalid."""
    try:
        y, w = divmod(int(yw), 100)
        d = datetime.date.fromisocalendar(y, w, 1) - datetime.timedelta(weeks=weeks)
        iy, iw, _ = d.isocalendar()
        return iy * 100 + iw
    except (ValueError, TypeError):
        return None


def _fb_values(telkomsel, isat3, xl, smartfren):
    """Map raw columns to (tsel, ioh, xls) where XLS = XL + Smartfren."""
    tsel = float(telkomsel or 0)
    ioh = float(isat3 or 0)
    xls = float(xl or 0) + float(smartfren or 0)
    return tsel, ioh, xls


def _avg_nonzero(values):
    """Average of values excluding 0/None (Excel AVERAGEIFS(..., "<>0")). Unrounded."""
    nz = [v for v in values if v]
    return (sum(nz) / len(nz)) if nz else None


def _r2(v):
    return round(v, 2) if v is not None else None


def _delta(cur_v, prev_v):
    if cur_v is None or prev_v is None:
        return None
    return round(cur_v - prev_v, 2)


def _fb_is_zero(tsel, ioh, xls):
    return not tsel and not ioh and not xls


def _fb_status(tsel, ioh, xls, win_lose):
    """DB win_lose has priority (a 0-value city flagged WIN is counted as WIN).
    Fallback when win_lose is empty: '-' for all-zero, else computed."""
    wl = (win_lose or "").strip().upper()
    if wl in ("WIN", "LOSE"):
        return wl
    if _fb_is_zero(tsel, ioh, xls):
        return "-"
    return "WIN" if tsel > max(ioh, xls) else "LOSE"


def _fb_olo(ioh, xls):
    """Strongest competitor (other licensed operator)."""
    return "XLS" if xls > ioh else "IOH"


FB_CITY_MASTER_START_YW = 202401
_FB_CITY_MASTER_CACHE = {"data": None, "timestamp": 0}


def get_cached_fb_city_master():
    """Distinct cities in sow.fb_share since 202401 with their latest branch: [(branch, city), ...]."""
    now = time.time()
    if _FB_CITY_MASTER_CACHE["data"] and (now - _FB_CITY_MASTER_CACHE["timestamp"] < _CACHE_TTL):
        return _FB_CITY_MASTER_CACHE["data"]
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(
                """
                SELECT branch, kabupaten FROM (
                    SELECT DISTINCT ON (kabupaten) kabupaten, branch
                    FROM sow.fb_share
                    WHERE yearweek >= %s AND kabupaten IS NOT NULL AND kabupaten <> ''
                    ORDER BY kabupaten, yearweek DESC
                ) t
                ORDER BY branch, kabupaten
                """,
                (FB_CITY_MASTER_START_YW,)
            )
            data = [(r[0], r[1]) for r in cur.fetchall()]
            _FB_CITY_MASTER_CACHE["data"] = data
            _FB_CITY_MASTER_CACHE["timestamp"] = now
            return data
    except Exception as e:
        logger.exception("Error loading fb_share city master: %s", e)
        return []


@sow.route("/api/sow/fb-share/trend", methods=["POST"])
@csrf.exempt
@login_required
def api_fb_share_trend():
    """FB Share trends: win-city achievement, regional/branch averages (non-zero) and city series."""
    req = request.get_json(silent=True) or {}
    try:
        start_yw = int(req.get("start_yearweek", 202501))
        end_yw = int(req.get("end_yearweek", 202652))
        if start_yw > end_yw:
            start_yw, end_yw = end_yw, start_yw
    except (ValueError, TypeError):
        return jsonify({"status": "error", "message": "Invalid yearweek range"}), 400

    query = """
        SELECT yearweek, region, branch, kabupaten, telkomsel, isat3, xl, smartfren, win_lose
        FROM sow.fb_share
        WHERE yearweek >= %s AND yearweek <= %s
        ORDER BY yearweek ASC, branch ASC, kabupaten ASC
    """
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(query, (start_yw, end_yw))
            rows = cur.fetchall()

        weeks = sorted({int(r[0]) for r in rows})
        w_idx = {w: i for i, w in enumerate(weeks)}
        n = len(weeks)

        regions = sorted({r[1] for r in rows if r[1]})
        master = get_cached_fb_city_master()
        branches = sorted({r[2] for r in rows if r[2]} | {b for b, _ in master if b})

        # Buckets for averaging: bucket[key][provider][week_idx] -> list of values
        def new_bucket():
            return {p: [[] for _ in range(n)] for p in FB_PROVIDERS}

        reg_bucket = new_bucket()
        br_buckets = {b: new_bucket() for b in branches}
        # City list seeded from master (distinct cities since 202401)
        city_map = {
            kab: {"city": kab, "branch": b, "tsel": [None] * n, "ioh": [None] * n, "xls": [None] * n}
            for b, kab in master
        }

        win_cnt = [0] * n
        zero_cnt = [0] * n
        total_cnt = [0] * n

        for r in rows:
            yw, _region, branch, kab = int(r[0]), r[1], r[2], r[3]
            i = w_idx[yw]
            tsel, ioh, xls = _fb_values(r[4], r[5], r[6], r[7])
            status = _fb_status(tsel, ioh, xls, r[8])

            total_cnt[i] += 1
            if _fb_is_zero(tsel, ioh, xls):
                zero_cnt[i] += 1
            # DB status has priority: 0-value city flagged WIN still counts as #win city
            if status == "WIN":
                win_cnt[i] += 1

            for p, v in zip(FB_PROVIDERS, (tsel, ioh, xls)):
                reg_bucket[p][i].append(v)
                if branch in br_buckets:
                    br_buckets[branch][p][i].append(v)

            if kab:
                c = city_map.get(kab)
                if c is None:
                    c = {"city": kab, "branch": branch, "tsel": [None] * n, "ioh": [None] * n, "xls": [None] * n}
                    city_map[kab] = c
                # 0 means no data -> keep as gap (None) in trend lines
                c["tsel"][i] = round(tsel, 2) if tsel else None
                c["ioh"][i] = round(ioh, 2) if ioh else None
                c["xls"][i] = round(xls, 2) if xls else None

        def finalize(bucket):
            return {p: [_r2(_avg_nonzero(vals)) for vals in bucket[p]] for p in FB_PROVIDERS}

        win_pct = [round(win_cnt[i] / total_cnt[i] * 100, 2) if total_cnt[i] else None for i in range(n)]

        cities = sorted(city_map.values(), key=lambda c: ((c["branch"] or ""), c["city"]))

        return jsonify({
            "status": "success",
            "start_yearweek": start_yw,
            "end_yearweek": end_yw,
            "weeks": [str(w) for w in weeks],
            "region_label": " / ".join(regions) if regions else "Regional",
            "branches": branches,
            "achievement": {
                "win": win_cnt,
                "zero": zero_cnt,
                "total": total_cnt,
                "win_pct": win_pct,
            },
            "regional": finalize(reg_bucket),
            "branch_series": {b: finalize(br_buckets[b]) for b in branches},
            "cities": cities,
        })
    except Exception as e:
        logger.exception("Error loading FB share trend: %s", e)
        return jsonify({"status": "error", "message": "Failed to fetch FB share trend"}), 500


@sow.route("/api/sow/fb-share/compare", methods=["POST"])
@csrf.exempt
@login_required
def api_fb_share_compare():
    """FB Share comparison tables for a reference week vs WoW (-1w), MoM (-4w), YoY (-1y)."""
    req = request.get_json(silent=True) or {}
    try:
        ref_yw = int(req.get("yearweek"))
    except (ValueError, TypeError):
        yws = get_cached_fb_yearweeks()
        ref_yw = yws[0] if yws else None
    if not ref_yw:
        return jsonify({"status": "error", "message": "No yearweek available"}), 400

    try:
        threshold = float(req.get("threshold", 95))
    except (ValueError, TypeError):
        threshold = 95.0

    cmp_weeks = {
        "wow": _shift_yearweek(ref_yw, 1),
        "mom": _shift_yearweek(ref_yw, 4),
        "yoy": ref_yw - 100,
    }
    all_weeks = [ref_yw] + [w for w in cmp_weeks.values() if w]

    query = """
        SELECT yearweek, region, branch, kabupaten, telkomsel, isat3, xl, smartfren, win_lose
        FROM sow.fb_share
        WHERE yearweek = ANY(%s)
    """
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(query, (all_weeks,))
            rows = cur.fetchall()

        # data[yw][kab] = dict(...)
        data = {}
        regions = set()
        for r in rows:
            yw = int(r[0])
            tsel, ioh, xls = _fb_values(r[4], r[5], r[6], r[7])
            if r[1]:
                regions.add(r[1])
            data.setdefault(yw, {})[r[3]] = {
                "branch": r[2], "tsel": tsel, "ioh": ioh, "xls": xls, "win_lose": r[8]
            }

        ref_data = data.get(ref_yw, {})
        if not ref_data:
            return jsonify({"status": "error", "message": f"No FB share data for week {ref_yw}"}), 404

        available = {k: (w if w in data else None) for k, w in cmp_weeks.items()}

        # ── Branch / Regional comparison (AVERAGEIFS <> 0) ──
        def group_avg(yw, branch=None):
            d = data.get(yw)
            if not d:
                return {p: None for p in FB_PROVIDERS}
            items = [v for v in d.values() if branch is None or v["branch"] == branch]
            return {p: _avg_nonzero([it[p] for it in items]) for p in FB_PROVIDERS}

        branches = sorted({v["branch"] for v in ref_data.values() if v["branch"]})
        region_label = " / ".join(sorted(regions)) if regions else "Regional"

        branch_rows = []
        for b in branches + [None]:
            cur_avg = group_avg(ref_yw, b)
            row = {
                "name": b if b else region_label,
                "is_regional": b is None,
                "share": {p: _r2(cur_avg[p]) for p in FB_PROVIDERS},
            }
            for k, w in available.items():
                prev = group_avg(w, b) if w else {p: None for p in FB_PROVIDERS}
                row[k] = {p: _delta(cur_avg[p], prev[p]) for p in FB_PROVIDERS}
            branch_rows.append(row)

        # ── City level table (city list = master since 202401 ∪ ref week) ──
        master = get_cached_fb_city_master()
        city_branch = {kab: b for b, kab in master}
        for kab, v in ref_data.items():
            city_branch.setdefault(kab, v["branch"])

        city_rows = []
        for kab, master_branch in city_branch.items():
            v = ref_data.get(kab)
            if v is None:
                # City exists in master list but has no record in the selected week
                row = {
                    "branch": master_branch, "city": kab,
                    "tsel": None, "ioh": None, "xls": None,
                    "status": "N/A", "olo": None, "zero": False, "missing": True,
                }
                for k in available:
                    row[k] = {p: None for p in FB_PROVIDERS}
                city_rows.append(row)
                continue
            status = _fb_status(v["tsel"], v["ioh"], v["xls"], v["win_lose"])
            row = {
                "branch": v["branch"] or master_branch,
                "city": kab,
                "tsel": round(v["tsel"], 2),
                "ioh": round(v["ioh"], 2),
                "xls": round(v["xls"], 2),
                "status": status,
                "olo": _fb_olo(v["ioh"], v["xls"]),
                "zero": _fb_is_zero(v["tsel"], v["ioh"], v["xls"]),
                "missing": False,
            }
            for k, w in available.items():
                prev = data.get(w, {}).get(kab) if w else None
                row[k] = {p: (_delta(v[p], prev[p]) if prev else None) for p in FB_PROVIDERS}
            city_rows.append(row)
        city_rows.sort(key=lambda c: ((c["branch"] or ""), c["city"]))

        # ── Summary ──
        regional = branch_rows[-1]
        reported = [c for c in city_rows if not c["missing"]]
        wins = [c for c in reported if c["status"] == "WIN"]  # includes 0-value cities flagged WIN
        loses = [c for c in reported if c["status"] == "LOSE"]
        zero = [c for c in reported if c["zero"]]
        below = [c for c in reported if c["tsel"] and c["tsel"] < threshold]
        lowest_city = min((c for c in reported if c["tsel"]), key=lambda c: c["tsel"], default=None)
        branch_only = [b for b in branch_rows if not b["is_regional"] and b["share"]["tsel"] is not None]
        best_branch = max(branch_only, key=lambda b: b["share"]["tsel"], default=None)
        low_branch = min(branch_only, key=lambda b: b["share"]["tsel"], default=None)
        r_share = regional["share"]
        olo_max = max([x for x in (r_share["ioh"], r_share["xls"]) if x is not None], default=None)

        summary = {
            "regional_tsel": r_share["tsel"],
            "regional_ioh": r_share["ioh"],
            "regional_xls": r_share["xls"],
            "regional_wow": regional["wow"]["tsel"],
            "regional_mom": regional["mom"]["tsel"],
            "regional_yoy": regional["yoy"]["tsel"],
            "gap_vs_olo": _delta(r_share["tsel"], olo_max),
            "win_city": len(wins),
            "lose_city": len(loses),
            "reported_city": len(reported),
            "master_city": len(city_rows),
            "missing_city": len(city_rows) - len(reported),
            "zero_city": len(zero),
            "win_rate": round(len(wins) / len(reported) * 100, 2) if reported else None,
            "below_threshold": len(below),
            "threshold": threshold,
            "lowest_city": {"city": lowest_city["city"], "branch": lowest_city["branch"], "tsel": lowest_city["tsel"]} if lowest_city else None,
            "best_branch": {"name": best_branch["name"], "tsel": best_branch["share"]["tsel"]} if best_branch else None,
            "low_branch": {"name": low_branch["name"], "tsel": low_branch["share"]["tsel"]} if low_branch else None,
        }

        return jsonify({
            "status": "success",
            "yearweek": ref_yw,
            "compare_weeks": cmp_weeks,
            "compare_available": {k: bool(w) for k, w in available.items()},
            "region_label": region_label,
            "branch_table": branch_rows,
            "city_table": city_rows,
            "summary": summary,
        })
    except Exception as e:
        logger.exception("Error loading FB share comparison: %s", e)
        return jsonify({"status": "error", "message": "Failed to fetch FB share comparison"}), 500


# ══════════════════════════════════════════════════════════════════════
# SOW TARGET SUMMARY (sow.sow_target)
# ══════════════════════════════════════════════════════════════════════
SUMMARY_SOW_TITLES = {
    "redcov": "Red Coverage",
    "unbalance": "Unbalance PRB",
    "rci": "RCI",
    "4g_good_thp": "4G Good Throughput",
    "5g_good_thp": "5G Good Throughput",
    "Payload": "Payload",
    "ONX": "ONX",
    "Ookla": "Ookla",
    "RHI": "RHI",
    "CEI": "CEI",
}

# SOW metrics whose quarterly targets are counts (not percentages)
SUMMARY_COUNT_METRICS = {"ONX", "Ookla"}


def _summary_title(key):
    return SUMMARY_SOW_TITLES.get(key, key)


def get_cached_summary_years():
    """Distinct years available in sow.sow_target, descending."""
    now = time.time()
    if _SUMMARY_YEARS_CACHE["data"] is not None and (now - _SUMMARY_YEARS_CACHE["timestamp"] < _CACHE_TTL):
        return _SUMMARY_YEARS_CACHE["data"]
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(
                'SELECT DISTINCT "Year" FROM sow.sow_target WHERE "Year" IS NOT NULL ORDER BY "Year" DESC'
            )
            years = [int(r[0]) for r in cur.fetchall()]
            _SUMMARY_YEARS_CACHE["data"] = years
            _SUMMARY_YEARS_CACHE["timestamp"] = now
            return years
    except Exception as e:
        logger.exception("Error loading sow_target years: %s", e)
        return []


@sow.route("/sow/summary")
@login_required
@viewer_blocked
def summary_page():
    """Render the SOW Summary page: quarterly targets per SOW metric for a chosen year."""
    years = get_cached_summary_years()
    default_year = years[0] if years else 2026

    return render_template(
        "sow_summary.html",
        username=session.get("username", "User"),
        years=years,
        default_year=default_year,
    )


@sow.route("/api/sow/targets", methods=["GET"])
@login_required
def api_sow_targets():
    """Return sow.sow_target rows (Baseline + Q1..Q4) filtered by year and optional Region."""
    try:
        year = int(request.args.get("year", 0))
    except (ValueError, TypeError):
        return jsonify({"status": "error", "message": "Invalid year"}), 400

    if not year:
        years = get_cached_summary_years()
        year = years[0] if years else None
    if not year:
        return jsonify({"status": "error", "message": "No year available in sow.sow_target"}), 400

    region = (request.args.get("region") or "").strip() or None

    query = """
        SELECT "SOW", "Baseline", "Q1", "Q2", "Q3", "Q4"
        FROM sow.sow_target
        WHERE "Year" = %s
    """
    params = [year]
    if region:
        query += ' AND "Region" = %s'
        params.append(region)

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(query, tuple(params))
            rows = cur.fetchall()

        items = []
        for r in rows:
            key = r[0]
            vals = [r[i] for i in range(1, 6)]
            items.append({
                "sow": key,
                "title": _summary_title(key),
                "is_count": key in SUMMARY_COUNT_METRICS,
                "baseline": vals[0],
                "q1": vals[1],
                "q2": vals[2],
                "q3": vals[3],
                "q4": vals[4],
            })
        items.sort(key=lambda x: x["title"].lower())

        return jsonify({
            "status": "success",
            "year": year,
            "region": region,
            "count": len(items),
            "targets": items,
        })
    except Exception as e:
        logger.exception("Error loading sow targets: %s", e)
        return jsonify({"status": "error", "message": "Failed to fetch SOW targets"}), 500


@sow.route("/sow/crowdsource")
@login_required
@viewer_blocked
def crowdsource_page():
    """Render the Crowdsource Scope of Work (SOW) page with multiple tabs."""
    yearweeks = get_cached_yearweeks()
    levels = ["Nation", "Area", "Region", "Kabupaten"]
    providers = ["Telkomsel", "Indosat", "XL", "Smartfren", "3"]
    kab_regions = get_cached_kab_regions()

    default_level = "Region"
    locations = get_cached_locations(default_level)
    default_location = "MALUKU DAN PAPUA" if "MALUKU DAN PAPUA" in locations else (locations[0] if locations else "")
    default_provider = "Telkomsel"

    default_start_yw = 202601 if 202601 in yearweeks else (yearweeks[-1] if yearweeks else 202601)
    default_end_yw = yearweeks[0] if yearweeks else 202638

    # FB Share tab defaults (independent week list from sow.fb_share)
    fb_yearweeks = get_cached_fb_yearweeks()
    fb_default_end_yw = fb_yearweeks[0] if fb_yearweeks else default_end_yw
    fb_default_start_yw = _shift_yearweek(fb_default_end_yw, 52) if fb_yearweeks else default_start_yw
    if fb_yearweeks and fb_default_start_yw not in fb_yearweeks:
        older = [y for y in fb_yearweeks if y <= (fb_default_start_yw or 0)]
        fb_default_start_yw = older[0] if older else fb_yearweeks[-1]

    return render_template(
        "sow_crowdsource.html",
        username=session.get("username", "User"),
        levels=levels,
        default_level=default_level,
        locations=locations,
        default_location=default_location,
        providers=providers,
        default_provider=default_provider,
        yearweeks=yearweeks,
        default_start_yw=default_start_yw,
        default_end_yw=default_end_yw,
        kab_regions=kab_regions,
        ordered_metrics=ORDERED_16_METRICS,
        metric_titles=METRIC_TITLES,
        fb_yearweeks=fb_yearweeks,
        fb_default_start_yw=fb_default_start_yw,
        fb_default_end_yw=fb_default_end_yw
    )


@sow.route("/api/sow/regions", methods=["GET"])
@login_required
def api_sow_regions():
    """Return distinct Region list for level='Kabupaten'."""
    regions = get_cached_kab_regions()
    return jsonify({"regions": regions})


@sow.route("/api/sow/locations", methods=["GET"])
@login_required
def api_sow_locations():
    """Return distinct locations for the given level and optional region."""
    level = request.args.get("level", "Region")
    region = request.args.get("region")
    locs = get_cached_locations(level, region)
    return jsonify({"level": level, "region": region, "locations": locs})


@sow.route("/api/sow/chart-data", methods=["POST"])
@csrf.exempt
@login_required
def api_sow_chart_data():
    """Return stacked bar chart data (#win and #lose) with quarterly custom targets supporting multi-year."""
    req = request.get_json(silent=True) or {}
    level = req.get("level", "Region")
    location = req.get("location", "MALUKU DAN PAPUA")
    provider = req.get("provider", "Telkomsel")

    try:
        start_yw = int(req.get("start_yearweek", 202601))
        end_yw = int(req.get("end_yearweek", 202638))
        if start_yw > end_yw:
            start_yw, end_yw = end_yw, start_yw
    except (ValueError, TypeError):
        start_yw = 202601
        end_yw = 202638

    # Quarterly targets from sow.sow_target (SOW = 'ONX'); fallback to 14.
    # Keyed by year so multi-year trends resolve the correct quarter.
    targets_by_year = {}
    target_q1 = 14
    target_q2 = 14
    target_q3 = 14
    target_q4 = 14
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute('SELECT "Year", "Q1", "Q2", "Q3", "Q4" FROM sow.sow_target WHERE "SOW" = %s', ("ONX",))
            for yr_, *qs in cur.fetchall():
                if yr_ is None:
                    continue
                vals = [int(q) if q is not None else 14 for q in qs]
                targets_by_year[int(yr_)] = dict(zip((1, 2, 3, 4), vals))
            if targets_by_year:
                latest = targets_by_year[max(targets_by_year)]
                target_q1, target_q2, target_q3, target_q4 = (latest[1], latest[2], latest[3], latest[4])
    except Exception as e:
        logger.exception("Error loading ONX targets from sow_target: %s", e)

    fallback_map = {1: target_q1, 2: target_q2, 3: target_q3, 4: target_q4}

    def _target_for(year, quarter):
        """Resolve ONX quarterly target for a year, falling back to the nearest
        year present in sow.sow_target (the table only carries certain years)."""
        year_map = targets_by_year.get(year)
        if not year_map and targets_by_year:
            nearest = min(targets_by_year, key=lambda y: (abs(y - year), y))
            year_map = targets_by_year[nearest]
        if year_map:
            return year_map.get(quarter, 14)
        return fallback_map.get(quarter, 14)


    query = """
        SELECT yearweek,
               SUM(CASE WHEN rank = 1 THEN 1 ELSE 0 END)::integer AS win_count,
               (16 - SUM(CASE WHEN rank = 1 THEN 1 ELSE 0 END))::integer AS lose_count
        FROM sow.onx_rank_score
        WHERE level = %s
          AND location = %s
          AND provider = %s
          AND yearweek >= %s
          AND yearweek <= %s
        GROUP BY yearweek
        ORDER BY yearweek ASC
    """

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(query, (level, location, provider, start_yw, end_yw))
            rows = cur.fetchall()

        labels = []
        win_data = []
        lose_data = []
        target_series = []

        target_met_count = 0
        total_wins = 0

        best_week = None
        max_win = -1
        lowest_week = None
        min_win = 999

        for r in rows:
            yw = int(r[0])
            w = int(r[1])
            l = int(r[2])
            if l < 0:
                l = 0

            # Determine target for this week's year and quarter (from sow.sow_target)
            yr = yw // 100
            week_num = yw % 100
            q = min(4, max(1, (week_num - 1) // 13 + 1))

            curr_target = _target_for(yr, q)

            labels.append(str(yw))
            win_data.append(w)
            lose_data.append(l)
            target_series.append(curr_target)

            total_wins += w
            if w >= curr_target:
                target_met_count += 1

            if w > max_win:
                max_win = w
                best_week = {"yearweek": yw, "win": w, "lose": l}
            if w < min_win:
                min_win = w
                lowest_week = {"yearweek": yw, "win": w, "lose": l}

        total_weeks = len(rows)
        latest_yw = rows[-1][0] if rows else None
        latest_win = rows[-1][1] if rows else 0
        latest_lose = rows[-1][2] if rows else 0
        latest_target = target_series[-1] if target_series else target_q1
        latest_rate = round((latest_win / 16.0) * 100, 1) if latest_win else 0.0

        avg_win = round(total_wins / total_weeks, 1) if total_weeks > 0 else 0.0
        avg_rate = round((avg_win / 16.0) * 100, 1) if total_weeks > 0 else 0.0
        target_met_pct = round((target_met_count / total_weeks) * 100, 1) if total_weeks > 0 else 0.0

        summary = {
            "total_weeks": total_weeks,
            "latest_yearweek": latest_yw,
            "latest_win": latest_win,
            "latest_lose": latest_lose,
            "latest_target": latest_target,
            "latest_rate": latest_rate,
            "target_q1": target_q1,
            "target_q2": target_q2,
            "target_q3": target_q3,
            "target_q4": target_q4,
            "target_met_count": target_met_count,
            "target_met_pct": target_met_pct,
            "avg_win": avg_win,
            "avg_rate": avg_rate,
            "best_week": best_week,
            "lowest_week": lowest_week,
        }

        return jsonify({
            "status": "success",
            "level": level,
            "location": location,
            "provider": provider,
            "start_yearweek": start_yw,
            "end_yearweek": end_yw,
            "labels": labels,
            "win": win_data,
            "lose": lose_data,
            "targets": target_series,
            "summary": summary
        })

    except Exception as e:
        logger.exception("Error executing chart query for SOW Crowdsource: %s", e)
        return jsonify({"status": "error", "message": "Failed to fetch chart data"}), 500


@sow.route("/api/sow/trend-metrics", methods=["POST"])
@csrf.exempt
@login_required
def api_sow_trend_metrics():
    """Return trend line scores and status rank winners for 16 metrics across multiple providers."""
    req = request.get_json(silent=True) or {}
    level = req.get("level", "Region")
    location = req.get("location", "MALUKU DAN PAPUA")
    providers = req.get("providers", ["Indosat", "Telkomsel", "XL"])
    if not isinstance(providers, list) or len(providers) == 0:
        providers = ["Indosat", "Telkomsel", "XL"]

    try:
        start_yw = int(req.get("start_yearweek", 202602))
        end_yw = int(req.get("end_yearweek", 202638))
        if start_yw > end_yw:
            start_yw, end_yw = end_yw, start_yw
    except (ValueError, TypeError):
        start_yw = 202602
        end_yw = 202638

    # Query 1: Metric data for the requested providers
    data_query = """
        SELECT yearweek, metric, provider, "mean or percentage" AS score, rank
        FROM sow.onx_rank_score
        WHERE level = %s
          AND location = %s
          AND provider = ANY(%s)
          AND yearweek >= %s
          AND yearweek <= %s
        ORDER BY yearweek ASC, metric ASC, provider ASC
    """

    # Query 2: Winners (rank = 1) across ALL providers in that location
    winners_query = """
        SELECT yearweek, metric, provider
        FROM sow.onx_rank_score
        WHERE level = %s
          AND location = %s
          AND rank = 1
          AND yearweek >= %s
          AND yearweek <= %s
        ORDER BY yearweek ASC, metric ASC, provider ASC
    """

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(data_query, (level, location, providers, start_yw, end_yw))
            rows = cur.fetchall()

            cur.execute(winners_query, (level, location, start_yw, end_yw))
            winner_rows = cur.fetchall()

        # Build list of weeks
        weeks_set = set()
        # Structure: metrics_data[metric][provider][yearweek] = {score, rank}
        metrics_data = {m: {p: {} for p in providers} for m in ORDERED_16_METRICS}

        for r in rows:
            yw = int(r[0])
            m = r[1]
            p = r[2]
            score = round(r[3], 2) if r[3] is not None else None
            rnk = r[4]

            weeks_set.add(yw)
            if m in metrics_data and p in metrics_data[m]:
                metrics_data[m][p][yw] = {"score": score, "rank": rnk}

        weeks_sorted = sorted(list(weeks_set))
        weeks_labels = [str(w) for w in weeks_sorted]

        # Winners structure: winners[metric][yearweek] = [providers with rank 1]
        winners_map = {m: {str(w): [] for w in weeks_sorted} for m in ORDERED_16_METRICS}
        for wr in winner_rows:
            yw_str = str(wr[0])
            m = wr[1]
            p = wr[2]
            if m in winners_map and yw_str in winners_map[m]:
                winners_map[m][yw_str].append(p)

        # Format metrics series aligned with weeks_sorted
        formatted_metrics = {}
        for m in ORDERED_16_METRICS:
            formatted_metrics[m] = {
                "title": METRIC_TITLES.get(m, m.replace("_", " ").title()),
                "providers": {}
            }
            for p in providers:
                score_series = []
                rank_series = []
                for w in weeks_sorted:
                    item = metrics_data[m][p].get(w)
                    if item:
                        score_series.append(item["score"])
                        rank_series.append(item["rank"])
                    else:
                        score_series.append(None)
                        rank_series.append(None)
                formatted_metrics[m]["providers"][p] = {
                    "scores": score_series,
                    "ranks": rank_series
                }

        return jsonify({
            "status": "success",
            "level": level,
            "location": location,
            "providers": providers,
            "weeks": weeks_labels,
            "ordered_metrics": ORDERED_16_METRICS,
            "metrics": formatted_metrics,
            "winners": winners_map
        })

    except Exception as e:
        logger.exception("Error loading trend metrics: %s", e)
        return jsonify({"status": "error", "message": "Failed to fetch trend metrics"}), 500


@sow.route("/api/sow/metrics-detail", methods=["GET"])
@login_required
def api_sow_metrics_detail():
    """Return the 16 individual metrics and their rank/status for a single week."""
    level = request.args.get("level", "Region")
    location = request.args.get("location", "MALUKU DAN PAPUA")
    provider = request.args.get("provider", "Telkomsel")
    try:
        yearweek = int(request.args.get("yearweek", 202638))
    except (ValueError, TypeError):
        yearweek = 202638

    query = """
        SELECT metric,
               rank,
               "mean or percentage" AS score,
               lci,
               uci,
               CASE 
                   WHEN rank IS NULL THEN 'no rank'
                   WHEN rank = 1 THEN 'win'
                   ELSE 'lose'
               END AS status,
               snapshot_date,
               week_date
        FROM sow.onx_rank_score
        WHERE level = %s
          AND location = %s
          AND provider = %s
          AND yearweek = %s
        ORDER BY 
          CASE WHEN rank = 1 THEN 1 WHEN rank IS NOT NULL THEN 2 ELSE 3 END,
          metric ASC
    """

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(query, (level, location, provider, yearweek))
            rows = cur.fetchall()

        metrics_list = []
        win_count = 0
        lose_count = 0
        no_rank_count = 0

        for r in rows:
            raw_metric = r[0]
            rank_val = r[1]
            score_val = round(r[2], 2) if r[2] is not None else None
            lci_val = round(r[3], 2) if r[3] is not None else None
            uci_val = round(r[4], 2) if r[4] is not None else None
            status_val = r[5]

            if status_val == "win":
                win_count += 1
            elif status_val == "lose":
                lose_count += 1
            else:
                no_rank_count += 1

            metrics_list.append({
                "metric": raw_metric,
                "metric_title": METRIC_TITLES.get(raw_metric, raw_metric.replace("_", " ").title()),
                "rank": rank_val,
                "score": score_val,
                "lci": lci_val,
                "uci": uci_val,
                "status": status_val
            })

        return jsonify({
            "status": "success",
            "yearweek": yearweek,
            "level": level,
            "location": location,
            "provider": provider,
            "total_metrics": len(metrics_list),
            "win_count": win_count,
            "lose_count": lose_count,
            "no_rank_count": no_rank_count,
            "metrics": metrics_list
        })
    except Exception as e:
        logger.exception("Error fetching metrics detail: %s", e)
        return jsonify({"status": "error", "message": "Failed to fetch metrics detail"}), 500


@sow.route("/api/sow/sample-template", methods=["GET"])
@login_required
def api_sow_sample_template():
    """Download sample CSV or XLSX template for editing and importing into sow.onx_rank_score."""
    file_format = request.args.get("format", "xlsx").lower()

    fieldnames = [
        "Region", "level", "location", "provider", "yearweek", "metric",
        "rank", "mean or percentage", "lci", "uci", "snapshot_date", "week_date"
    ]
    sample_rows = []
    for m in ORDERED_16_METRICS:
        sample_rows.append([
            "MALUKU DAN PAPUA", "Region", "MALUKU DAN PAPUA", "Telkomsel",
            202639, m, 1, 50.0, 48.5, 51.5,
            "2026-09-28 00:00:00", "2026-09-28 00:00:00"
        ])

    if file_format == "csv":
        csv_buffer = io.StringIO()
        writer = csv.writer(csv_buffer)
        writer.writerow(fieldnames)
        writer.writerows(sample_rows)
        output = make_response(csv_buffer.getvalue())
        output.headers["Content-Disposition"] = "attachment; filename=SOW_ONX_Import_Sample.csv"
        output.headers["Content-type"] = "text/csv; charset=utf-8"
        return output
    else:
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "onx_rank_score"
            ws.append(fieldnames)
            for row in sample_rows:
                ws.append(row)

            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)
            return send_file(
                excel_buffer,
                download_name="SOW_ONX_Import_Sample.xlsx",
                as_attachment=True,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            logger.exception("Error generating Excel template: %s", e)
            return jsonify({"status": "error", "message": "Failed to generate Excel template"}), 500


@sow.route("/api/sow/import", methods=["POST"])
@csrf.exempt
@login_required
def api_sow_import():
    """Import and upsert records into sow.onx_rank_score from uploaded CSV or Excel file."""
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file uploaded"}), 400

    uploaded_file = request.files["file"]
    if not uploaded_file.filename:
        return jsonify({"status": "error", "message": "Empty file name"}), 400

    filename = uploaded_file.filename.lower()
    raw_dicts = []

    try:
        if filename.endswith(".csv"):
            file_bytes = uploaded_file.read()
            text = file_bytes.decode("utf-8-sig", errors="replace")
            stream = io.StringIO(text)
            reader = csv.DictReader(stream)
            raw_dicts = list(reader)
        elif filename.endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(uploaded_file.stream, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                return jsonify({"status": "error", "message": "Uploaded Excel sheet is empty"}), 400
            headers = [str(c).strip() if c is not None else "" for c in rows[0]]
            for r in rows[1:]:
                if any(v is not None for v in r):
                    row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                    raw_dicts.append(row_dict)
        else:
            return jsonify({"status": "error", "message": "Only CSV and Excel (.xlsx, .xls) files are supported"}), 400
    except Exception as e:
        logger.exception("Failed to parse uploaded file: %s", e)
        return jsonify({"status": "error", "message": f"File parsing failed: {str(e)}"}), 400

    if not raw_dicts:
        return jsonify({"status": "error", "message": "No data rows found in uploaded file"}), 400

    # Clean headers in row dictionaries
    cleaned_rows = []
    for r in raw_dicts:
        cleaned_rows.append({str(k).strip(): v for k, v in r.items() if k is not None})

    first_row = cleaned_rows[0]
    required_cols = ["level", "location", "provider", "yearweek", "metric"]
    missing = [c for c in required_cols if c not in first_row]
    if missing:
        return jsonify({
            "status": "error",
            "message": f"Missing required columns in file: {', '.join(missing)}"
        }), 400

    # Detect score column
    score_col = None
    for cand in ["mean or percentage", "score", "value", "mean_or_percentage"]:
        if cand in first_row:
            score_col = cand
            break

    if not score_col:
        return jsonify({
            "status": "error",
            "message": "Missing 'mean or percentage' score column in uploaded file"
        }), 400

    upsert_query = """
        INSERT INTO sow.onx_rank_score (
            "Region", level, location, provider, yearweek, metric,
            rank, "mean or percentage", lci, uci, snapshot_date, week_date
        ) VALUES (
            %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s, COALESCE(%s, CURRENT_TIMESTAMP), COALESCE(%s, CURRENT_TIMESTAMP)
        )
        ON CONFLICT (yearweek, metric, level, location, provider)
        DO UPDATE SET
            rank = EXCLUDED.rank,
            "mean or percentage" = EXCLUDED."mean or percentage",
            lci = COALESCE(EXCLUDED.lci, sow.onx_rank_score.lci),
            uci = COALESCE(EXCLUDED.uci, sow.onx_rank_score.uci),
            "Region" = COALESCE(EXCLUDED."Region", sow.onx_rank_score."Region"),
            snapshot_date = COALESCE(EXCLUDED.snapshot_date, sow.onx_rank_score.snapshot_date, CURRENT_TIMESTAMP),
            week_date = COALESCE(EXCLUDED.week_date, sow.onx_rank_score.week_date, CURRENT_TIMESTAMP)
    """

    records = []
    for row in cleaned_rows:
        region = str(row.get("Region") or "").strip()
        if not region or region in ("None", "nan"):
            region = "-"

        lvl = str(row.get("level") or "").strip()
        loc = str(row.get("location") or "").strip()
        prv = str(row.get("provider") or "").strip()

        try:
            yw = int(float(str(row.get("yearweek")).strip()))
        except (ValueError, TypeError):
            continue

        metric = str(row.get("metric") or "").strip()

        # Rank
        rank_val = None
        raw_rank = row.get("rank")
        if raw_rank is not None and str(raw_rank).strip() not in ("", "None", "nan"):
            try:
                rank_val = int(float(str(raw_rank).strip()))
            except (ValueError, TypeError):
                rank_val = None

        # Score
        score_val = None
        raw_score = row.get(score_col)
        if raw_score is not None and str(raw_score).strip() not in ("", "None", "nan"):
            try:
                score_val = float(str(raw_score).strip())
            except (ValueError, TypeError):
                score_val = None

        # LCI & UCI
        lci_val = None
        raw_lci = row.get("lci")
        if raw_lci is not None and str(raw_lci).strip() not in ("", "None", "nan"):
            try:
                lci_val = float(str(raw_lci).strip())
            except (ValueError, TypeError):
                lci_val = None

        uci_val = None
        raw_uci = row.get("uci")
        if raw_uci is not None and str(raw_uci).strip() not in ("", "None", "nan"):
            try:
                uci_val = float(str(raw_uci).strip())
            except (ValueError, TypeError):
                uci_val = None

        # snapshot_date & week_date
        snap_val = None
        raw_snap = row.get("snapshot_date")
        if raw_snap is not None and str(raw_snap).strip() not in ("", "None", "nan"):
            snap_val = str(raw_snap).strip()

        week_date_val = None
        raw_wk = row.get("week_date")
        if raw_wk is not None and str(raw_wk).strip() not in ("", "None", "nan"):
            week_date_val = str(raw_wk).strip()

        records.append((
            region, lvl, loc, prv, yw, metric,
            rank_val, score_val, lci_val, uci_val,
            snap_val, week_date_val
        ))

    if not records:
        return jsonify({"status": "error", "message": "No valid data rows found in uploaded file"}), 400

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            psycopg2.extras.execute_batch(cur, upsert_query, records, page_size=200)
            conn.commit()

        clear_sow_caches()

        return jsonify({
            "status": "success",
            "rows_processed": len(records),
            "message": f"Successfully imported and updated {len(records)} benchmark records."
        })
    except Exception as e:
        logger.exception("Database error while upserting SOW records: %s", e)
        return jsonify({"status": "error", "message": f"Database error during import: {str(e)}"}), 500


# ══════════════════════════════════════════════════════════════════════
# FB SHARE import / sample template (sow.fb_share)
# ══════════════════════════════════════════════════════════════════════
FB_SHARE_FIELDS = [
    "yearweek", "region", "branch", "kabupaten", "archetype",
    "telkomsel", "isat3", "xl", "smartfren",
    "gap", "wow", "mom", "wow_cc", "mom_cc", "win_lose",
]


def _fb_share_sample_row(region="MALUKU DAN PAPUA", branch="AMBON", city="KOTA AMBON", yw=202639):
    return [
        yw, region, branch, city, "Super Fortress",
        91.5, 6.4, 0.8, 0.2,
        85.1, 0.3, 1.2, 0.1, 0.9, "WIN",
    ]


@sow.route("/api/sow/fb-sample-template", methods=["GET"])
@login_required
def api_sow_fb_sample_template():
    """Download sample CSV or XLSX template for importing into sow.fb_share."""
    file_format = request.args.get("format", "xlsx").lower()
    fieldnames = FB_SHARE_FIELDS
    sample_rows = [
        _fb_share_sample_row("MALUKU DAN PAPUA", "AMBON", "KOTA AMBON", 202639),
        _fb_share_sample_row("MALUKU DAN PAPUA", "JAYAPURA", "KOTA JAYAPURA", 202639),
        _fb_share_sample_row("MALUKU DAN PAPUA", "SORONG", "KOTA SORONG", 202639),
    ]

    if file_format == "csv":
        csv_buffer = io.StringIO()
        writer = csv.writer(csv_buffer)
        writer.writerow(fieldnames)
        writer.writerows(sample_rows)
        output = make_response(csv_buffer.getvalue())
        output.headers["Content-Disposition"] = "attachment; filename=SOW_FBShare_Import_Sample.csv"
        output.headers["Content-type"] = "text/csv; charset=utf-8"
        return output
    else:
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "fb_share"
            ws.append(fieldnames)
            for row in sample_rows:
                ws.append(row)
            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)
            return send_file(
                excel_buffer,
                download_name="SOW_FBShare_Import_Sample.xlsx",
                as_attachment=True,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            logger.exception("Error generating FB Share Excel template: %s", e)
            return jsonify({"status": "error", "message": "Failed to generate Excel template"}), 500


@sow.route("/api/sow/fb-import", methods=["POST"])
@csrf.exempt
@login_required
def api_sow_fb_import():
    """Import and upsert records into sow.fb_share from uploaded CSV or Excel file."""
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file uploaded"}), 400

    uploaded_file = request.files["file"]
    if not uploaded_file.filename:
        return jsonify({"status": "error", "message": "Empty file name"}), 400

    filename = uploaded_file.filename.lower()
    raw_dicts = []

    try:
        if filename.endswith(".csv"):
            file_bytes = uploaded_file.read()
            text = file_bytes.decode("utf-8-sig", errors="replace")
            stream = io.StringIO(text)
            reader = csv.DictReader(stream)
            raw_dicts = list(reader)
        elif filename.endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(uploaded_file.stream, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                return jsonify({"status": "error", "message": "Uploaded Excel sheet is empty"}), 400
            headers = [str(c).strip() if c is not None else "" for c in rows[0]]
            for r in rows[1:]:
                if any(v is not None for v in r):
                    row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                    raw_dicts.append(row_dict)
        else:
            return jsonify({"status": "error", "message": "Only CSV and Excel (.xlsx, .xls) files are supported"}), 400
    except Exception as e:
        logger.exception("Failed to parse uploaded file: %s", e)
        return jsonify({"status": "error", "message": f"File parsing failed: {str(e)}"}), 400

    if not raw_dicts:
        return jsonify({"status": "error", "message": "No data rows found in uploaded file"}), 400

    cleaned_rows = [{str(k).strip(): v for k, v in r.items() if k is not None} for r in raw_dicts]

    first_row = cleaned_rows[0]
    required_cols = ["yearweek", "branch", "kabupaten"]
    missing = [c for c in required_cols if c not in first_row]
    if missing:
        return jsonify({"status": "error", "message": f"Missing required columns in file: {', '.join(missing)}"}), 400

    def _num(v):
        if v is None or str(v).strip() in ("", "None", "nan"):
            return None
        try:
            return float(str(v).strip().replace(",", "."))
        except (ValueError, TypeError):
            return None

    records = []
    for row in cleaned_rows:
        try:
            yw = int(float(str(row.get("yearweek")).strip()))
        except (ValueError, TypeError):
            continue
        branch = str(row.get("branch") or "").strip()
        kabupaten = str(row.get("kabupaten") or "").strip()
        if not branch or not kabupaten:
            continue
        region = str(row.get("region") or "").strip() or None
        archetype = str(row.get("archetype") or "").strip() or None
        win_lose = str(row.get("win_lose") or "").strip() or None
        records.append((
            yw, region, branch, kabupaten, archetype,
            _num(row.get("telkomsel")), _num(row.get("isat3")),
            _num(row.get("xl")), _num(row.get("smartfren")),
            _num(row.get("gap")), _num(row.get("wow")), _num(row.get("mom")),
            _num(row.get("wow_cc")), _num(row.get("mom_cc")), win_lose,
        ))

    if not records:
        return jsonify({"status": "error", "message": "No valid data rows found in uploaded file"}), 400

    upsert_query = """
        INSERT INTO sow.fb_share (
            yearweek, region, branch, kabupaten, archetype,
            telkomsel, isat3, xl, smartfren,
            gap, wow, mom, wow_cc, mom_cc, win_lose
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (yearweek, branch, kabupaten)
        DO UPDATE SET
            region = COALESCE(EXCLUDED.region, sow.fb_share.region),
            archetype = COALESCE(EXCLUDED.archetype, sow.fb_share.archetype),
            telkomsel = EXCLUDED.telkomsel,
            isat3 = EXCLUDED.isat3,
            xl = EXCLUDED.xl,
            smartfren = EXCLUDED.smartfren,
            gap = EXCLUDED.gap,
            wow = EXCLUDED.wow,
            mom = EXCLUDED.mom,
            wow_cc = EXCLUDED.wow_cc,
            mom_cc = EXCLUDED.mom_cc,
            win_lose = COALESCE(EXCLUDED.win_lose, sow.fb_share.win_lose)
    """

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            psycopg2.extras.execute_batch(cur, upsert_query, records, page_size=200)
            conn.commit()

        clear_sow_caches()

        return jsonify({
            "status": "success",
            "rows_processed": len(records),
            "message": f"Successfully imported and updated {len(records)} FB Share records."
        })
    except Exception as e:
        logger.exception("Database error while upserting FB Share records: %s", e)
        return jsonify({"status": "error", "message": f"Database error during import: {str(e)}"}), 500
