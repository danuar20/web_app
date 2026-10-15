"""RHI Dashboard Routes — /sow/rhi

Radio Health Index (RHI) dashboard based on sow.rhi_score and sow.sow_target.
Provides:
  1. Top KPI Summary & Regional Benchmark Gauge (Good cell PUMA + Branch strip distribution)
  2. Technology (2G/4G/5G) Health & Key RHI Indicator Cards with mini sparklines & WoW Delta
  3. Weekly Trend Combination Chart (Stacked Bar Good/Fail counts + % Good Cell Line + Target)
  4. Technology Trend & Top Fail Root Cause Breakdown (availability, cssr, ccsr, volte, etc.)
  5. Comparison Table: NOP Level WoW, Root Cause WoW, and Kabupaten Level WoW
  6. Sample template download and Data Import for sow.rhi_score
"""
import csv
import io
import logging
import os
import tempfile
import threading
import time
from collections import defaultdict
from uuid import uuid4
import psycopg2.extras
from flask import Blueprint, render_template, request, session, jsonify, make_response, send_file
from app import csrf
from app.db.db_webapp import get_postgres_connection
from ._utils import login_required, viewer_blocked, db_query, csv_response

logger = logging.getLogger(__name__)

rhi = Blueprint("rhi", __name__)

_RHI_CACHE = {}
_RHI_CACHE_TTL = 900  # 15 minutes

# In-memory job progress tracking for async background imports
_IMPORT_JOBS = {}
_IMPORT_JOBS_LOCK = threading.Lock()

# Top Fail KPI labels map
FAIL_REASON_LABELS = {
    "availability_4g": "4G Availability",
    "cssr_data_service": "4G Data Setup (CSSR)",
    "availability_2g": "2G Availability",
    "ccsr_voice_rate": "2G Voice Drop (CCSR)",
    "ccsr_data_service": "4G Data Drop (CCSR)",
    "volte_retain_rate": "VoLTE Retainability",
    "sd_setup_sr_rate": "2G SDCCH Setup",
    "volte_accessibility_sr": "VoLTE Accessibility",
    "cssr_voice_rate": "2G Voice Setup (CSSR)",
    "ccsr_data_service_5g": "5G Data Drop",
    "cssr_data_service_5g": "5G Data Setup",
}

# Key RHI Indicator definitions
RHI_INDICATORS = [
    {
        "id": "cssr_data",
        "title": "4G CSSR Data",
        "tech": "4g",
        "val_col": "cssr_data_service",
        "fail_col": "total_fail_cssr_data_service_remark",
        "color": "#2563eb",
        "unit": "%",
    },
    {
        "id": "ccsr_data",
        "title": "4G CCSR Data",
        "tech": "4g",
        "val_col": "ccsr_data_service",
        "fail_col": "total_fail_ccsr_data_service_remark",
        "color": "#4f46e5",
        "unit": "%",
    },
    {
        "id": "volte_acc",
        "title": "VoLTE Accessibility",
        "tech": "4g",
        "val_col": "volte_accessibility_sr",
        "fail_col": "total_fail_volte_accessibility_sr_remark",
        "color": "#059669",
        "unit": "%",
    },
    {
        "id": "volte_ret",
        "title": "VoLTE Retainability",
        "tech": "4g",
        "val_col": "volte_retain_rate",
        "fail_col": "total_fail_volte_retain_rate_remark",
        "color": "#d97706",
        "unit": "%",
    },
    {
        "id": "ccsr_2g",
        "title": "2G CCSR Voice",
        "tech": "2g",
        "val_col": "ccsr_tch",
        "fail_col": "total_fail_ccsr_tch_remark",
        "color": "#0284c7",
        "unit": "%",
    },
    {
        "id": "sdsr_2g",
        "title": "2G SDSR",
        "tech": "2g",
        "val_col": "sdsr",
        "fail_col": "total_fail_sdsr_remark",
        "color": "#7c3aed",
        "unit": "%",
    },
]

# All columns of sow.rhi_score for import & export
RHI_COLUMNS = [
    'weeknum', 'tech', 'regional', 'vendor', 'lac', 'ci', 'cell_name', 'site_id', 'band', 'clutter',
    'kabupaten', 'id_kabupaten', 'kecamatan', 'id_kecamatan', 'total_fail_daily', 'cssr_tch', 'sdsr',
    'ccsr_tch', 'cssr_tch_remark', 'sdsr_remark', 'ccsr_tch_remark', 'cssr_data_service', 'ccsr_data_service',
    'volte_accessibility_sr', 'volte_retain_rate', 'cssr_data_service_remark', 'ccsr_data_service_remark',
    'volte_accessibility_sr_remark', 'volte_retain_rate_remark', 'cssr_data_service_5g', 'ccsr_data_service_5g',
    'cssr_data_service_5g_remark', 'ccsr_data_service_5g_remark', 'accessibility_remark', 'retainability_remark',
    'total_fail_cssr_tch_remark', 'total_fail_sdsr_remark', 'total_fail_ccsr_tch_remark',
    'total_fail_cssr_data_service_remark', 'total_fail_ccsr_data_service_remark',
    'total_fail_volte_accessibility_sr_remark', 'total_fail_volte_retain_rate_remark',
    'total_fail_cssr_data_service_5g_remark', 'total_fail_ccsr_data_service_5g_remark', 'remark_week',
    'availability_site_rate', 'availability_site_remark', 'availability_remark',
    'total_fail_availability_site_remark', 'most_fail_kpi_2g', 'most_fail_kpi_4g', 'most_fail_kpi_5g',
    'cssr_tch_d1', 'cssr_tch_d2', 'cssr_tch_d3', 'cssr_tch_d4', 'cssr_tch_d5', 'cssr_tch_d6', 'cssr_tch_d7',
    'sdsr_d1', 'sdsr_d2', 'sdsr_d3', 'sdsr_d4', 'sdsr_d5', 'sdsr_d6', 'sdsr_d7', 'ccsr_tch_d1', 'ccsr_tch_d2',
    'ccsr_tch_d3', 'ccsr_tch_d4', 'ccsr_tch_d5', 'ccsr_tch_d6', 'ccsr_tch_d7', 'cssr_data_service_d1',
    'cssr_data_service_d2', 'cssr_data_service_d3', 'cssr_data_service_d4', 'cssr_data_service_d5',
    'cssr_data_service_d6', 'cssr_data_service_d7', 'ccsr_data_service_d1', 'ccsr_data_service_d2',
    'ccsr_data_service_d3', 'ccsr_data_service_d4', 'ccsr_data_service_d5', 'ccsr_data_service_d6',
    'ccsr_data_service_d7', 'volte_accessibility_sr_d1', 'volte_accessibility_sr_d2', 'volte_accessibility_sr_d3',
    'volte_accessibility_sr_d4', 'volte_accessibility_sr_d5', 'volte_accessibility_sr_d6', 'volte_accessibility_sr_d7',
    'volte_retain_rate_d1', 'volte_retain_rate_d2', 'volte_retain_rate_d3', 'volte_retain_rate_d4',
    'volte_retain_rate_d5', 'volte_retain_rate_d6', 'volte_retain_rate_d7', 'cssr_data_service_5g_d1',
    'cssr_data_service_5g_d2', 'cssr_data_service_5g_d3', 'cssr_data_service_5g_d4', 'cssr_data_service_5g_d5',
    'cssr_data_service_5g_d6', 'cssr_data_service_5g_d7', 'ccsr_data_service_5g_d1', 'ccsr_data_service_5g_d2',
    'ccsr_data_service_5g_d3', 'ccsr_data_service_5g_d4', 'ccsr_data_service_5g_d5', 'ccsr_data_service_5g_d6',
    'ccsr_data_service_5g_d7', 'availability_d1', 'availability_d2', 'availability_d3', 'availability_d4',
    'availability_d5', 'availability_d6', 'availability_d7', 'cssr_tch_fail_rate', 'sdsr_fail_rate',
    'ccsr_tch_fail_rate', 'cssr_data_service_fail_rate', 'ccsr_data_service_fail_rate',
    'volte_accessibility_sr_fail_rate', 'volte_retain_fail_rate', 'cssr_data_service_5g_fail_rate',
    'ccsr_data_service_5g_fail_rate', 'availability_site_fail_rate', 'desa', 'id_desa', 'region_new',
    'dl_bler_rate', 'ul_bler_rate', 'dl_bler_rate_d1', 'dl_bler_rate_d2', 'dl_bler_rate_d3', 'dl_bler_rate_d4',
    'dl_bler_rate_d5', 'dl_bler_rate_d6', 'dl_bler_rate_d7', 'ul_bler_rate_d1', 'ul_bler_rate_d2',
    'ul_bler_rate_d3', 'ul_bler_rate_d4', 'ul_bler_rate_d5', 'ul_bler_rate_d6', 'ul_bler_rate_d7',
    'dl_bler_rate_remark', 'ul_bler_rate_remark', 'total_fail_dl_bler_rate_remark',
    'total_fail_ul_bler_rate_remark', 'type', 'availability_cellbasedconsecutive_remark',
    'total_fail_availability_cellbasedconsecutive_remark', 'avail_cell_w1', 'avail_cell_w2', 'avail_cell_w3',
    'avail_cell_w4', 'red_site_bucket_avail', 'total_fail_availability_site_remark_simul5'
]

RHI_INT_COLS = {'weeknum', 'band', 'desa', 'id_desa', 'total_fail_ul_bler_rate_remark'}
RHI_FLOAT_COLS = {
    'lac', 'ci', 'id_kabupaten', 'id_kecamatan', 'total_fail_daily', 'cssr_tch', 'sdsr', 'ccsr_tch',
    'cssr_tch_remark', 'sdsr_remark', 'ccsr_tch_remark', 'cssr_data_service', 'ccsr_data_service',
    'volte_accessibility_sr', 'volte_retain_rate', 'cssr_data_service_remark', 'ccsr_data_service_remark',
    'volte_accessibility_sr_remark', 'volte_retain_rate_remark', 'cssr_data_service_5g', 'ccsr_data_service_5g',
    'cssr_data_service_5g_remark', 'ccsr_data_service_5g_remark', 'accessibility_remark', 'retainability_remark',
    'total_fail_cssr_tch_remark', 'total_fail_sdsr_remark', 'total_fail_ccsr_tch_remark',
    'total_fail_cssr_data_service_remark', 'total_fail_ccsr_data_service_remark',
    'total_fail_volte_accessibility_sr_remark', 'total_fail_volte_retain_rate_remark',
    'total_fail_cssr_data_service_5g_remark', 'total_fail_ccsr_data_service_5g_remark', 'cssr_tch_d1',
    'cssr_tch_d2', 'cssr_tch_d3', 'cssr_tch_d4', 'cssr_tch_d5', 'cssr_tch_d6', 'cssr_tch_d7', 'sdsr_d1',
    'sdsr_d2', 'sdsr_d3', 'sdsr_d4', 'sdsr_d5', 'sdsr_d6', 'sdsr_d7', 'ccsr_tch_d1', 'ccsr_tch_d2',
    'ccsr_tch_d3', 'ccsr_tch_d4', 'ccsr_tch_d5', 'ccsr_tch_d6', 'ccsr_tch_d7', 'cssr_data_service_d1',
    'cssr_data_service_d2', 'cssr_data_service_d3', 'cssr_data_service_d4', 'cssr_data_service_d5',
    'cssr_data_service_d6', 'cssr_data_service_d7', 'ccsr_data_service_d1', 'ccsr_data_service_d2',
    'ccsr_data_service_d3', 'ccsr_data_service_d4', 'ccsr_data_service_d5', 'ccsr_data_service_d6',
    'ccsr_data_service_d7', 'volte_accessibility_sr_d1', 'volte_accessibility_sr_d2', 'volte_accessibility_sr_d3',
    'volte_accessibility_sr_d4', 'volte_accessibility_sr_d5', 'volte_accessibility_sr_d6',
    'volte_accessibility_sr_d7', 'volte_retain_rate_d1', 'volte_retain_rate_d2', 'volte_retain_rate_d3',
    'volte_retain_rate_d4', 'volte_retain_rate_d5', 'volte_retain_rate_d6', 'volte_retain_rate_d7',
    'cssr_data_service_5g_d1', 'cssr_data_service_5g_d2', 'cssr_data_service_5g_d3', 'cssr_data_service_5g_d4',
    'cssr_data_service_5g_d5', 'cssr_data_service_5g_d6', 'cssr_data_service_5g_d7', 'ccsr_data_service_5g_d1',
    'ccsr_data_service_5g_d2', 'ccsr_data_service_5g_d3', 'ccsr_data_service_5g_d4', 'ccsr_data_service_5g_d5',
    'ccsr_data_service_5g_d6', 'ccsr_data_service_5g_d7', 'ul_bler_rate', 'ul_bler_rate_d1', 'ul_bler_rate_d2',
    'ul_bler_rate_d3', 'ul_bler_rate_d4', 'ul_bler_rate_d5', 'ul_bler_rate_d6', 'ul_bler_rate_d7',
    'ul_bler_rate_remark', 'total_fail_availability_cellbasedconsecutive_remark', 'avail_cell_w1',
    'avail_cell_w2', 'avail_cell_w3', 'avail_cell_w4', 'total_fail_availability_site_remark_simul5'
}


def clear_rhi_caches():
    _RHI_CACHE.clear()


def _safe_str(val, default=None):
    if val is None:
        return default
    s = str(val).strip()
    return s if s != "" else default


def _safe_int(val, default=0):
    if val is None or val == "":
        return default
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return default


def _safe_float(val, default=None):
    if val is None or val == "":
        return default
    try:
        return float(str(val).strip())
    except (ValueError, TypeError):
        return default


def _get_branch_map():
    now = time.time()
    if "_branch_map" in _RHI_CACHE:
        val, exp = _RHI_CACHE["_branch_map"]
        if now < exp:
            return val
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(
                "SELECT DISTINCT branch, UPPER(TRIM(kabupaten)) "
                "FROM sow.fb_share WHERE branch IS NOT NULL AND kabupaten IS NOT NULL"
            )
            mapping = {}
            for b, k in cur.fetchall():
                if k not in mapping:
                    mapping[k] = b
            # Aliases present in sow.rhi_score
            mapping["FAK FAK"] = "SORONG"
            mapping["KEPULAUAN TANIMBAR"] = "AMBON"
            _RHI_CACHE["_branch_map"] = (mapping, now + 86400)
            return mapping
    except Exception as e:
        logger.exception("Error loading branch map for RHI: %s", e)
        return {
            "FAK FAK": "SORONG",
            "KEPULAUAN TANIMBAR": "AMBON",
        }


def _get_yearweeks():
    """Get list of valid weeknums with actual RHI data, cached in memory."""
    now = time.time()
    if "_yearweeks" in _RHI_CACHE:
        val, exp = _RHI_CACHE["_yearweeks"]
        if now < exp:
            return val

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute("""
                SELECT weeknum
                FROM sow.rhi_score
                WHERE weeknum IS NOT NULL 
                  AND weeknum >= 202001 
                  AND weeknum <= 202952
                GROUP BY weeknum
                HAVING COUNT(*) >= 50
                   AND COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) IN ('GOOD', 'FAIL')) >= 50
                ORDER BY weeknum DESC
            """)
            yws = [int(r[0]) for r in cur.fetchall() if r[0] is not None]
            _RHI_CACHE["_yearweeks"] = (yws, now + 60)
            return yws
    except Exception as e:
        logger.exception("Error loading RHI yearweeks: %s", e)
        return []


def _quarter_for_week(yearweek):
    try:
        week = int(str(yearweek)[4:])
    except Exception:
        return "Q1"
    if week <= 13:
        return "Q1"
    if week <= 26:
        return "Q2"
    if week <= 39:
        return "Q3"
    return "Q4"


def _get_all_targets():
    """All RHI targets from sow.sow_target, keyed by Year -> {'Q1': pct, ...}."""
    now = time.time()
    cached = _RHI_CACHE.get("_all_targets")
    if cached and now < cached[1]:
        return cached[0]
    targets = {}
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(
                'SELECT "Year", "Q1", "Q2", "Q3", "Q4" '
                'FROM sow.sow_target WHERE UPPER(TRIM("SOW")) = \'RHI\''
            )
            for year, *qs in cur.fetchall():
                q_map = {}
                for name, raw in zip(("Q1", "Q2", "Q3", "Q4"), qs):
                    if raw is None:
                        q_map[name] = None
                        continue
                    v = float(raw)
                    if v <= 1.0:
                        v *= 100.0
                    q_map[name] = round(v, 2)
                targets[int(year)] = q_map
    except Exception as e:
        logger.warning("Error fetching RHI targets: %s", e)
    _RHI_CACHE["_all_targets"] = (targets, now + 3600)
    return targets


def _get_target_for_week(yearweek):
    if not yearweek:
        return 93.55, "Target Q3"
    year = int(str(yearweek)[:4])
    q_col = _quarter_for_week(yearweek)
    val = _get_all_targets().get(year, {}).get(q_col)
    if val is None:
        return 93.55, f"Target {q_col}"
    return val, f"Target {q_col}"


# ══════════════════════════════════════════════════════════════════════
# Page View
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/sow/rhi")
@rhi.route("/okr/rhi")
@login_required
def rhi_dashboard_page():
    if request.args.get("refresh") == "1":
        clear_rhi_caches()
    yearweeks = _get_yearweeks()
    default_yw = yearweeks[0] if yearweeks else 202639
    default_start_yw = yearweeks[-1] if yearweeks else 202637
    default_gauge_yw = yearweeks[0] if yearweeks else 202639
    latest_week_label = f"W{str(default_yw)[4:]} ({default_yw})" if default_yw else "-"
    target_val, target_label = _get_target_for_week(default_yw)
    baseline_val = 93.05

    return render_template(
        "sow_rhi.html",
        username=session.get("username", "User"),
        yearweeks=yearweeks,
        latest_week=latest_week_label,
        default_yw=default_yw,
        default_start_yw=default_start_yw,
        default_gauge_yw=default_gauge_yw,
        indicators=RHI_INDICATORS,
        target_val=target_val,
        target_label=target_label,
        baseline_val=baseline_val,
    )


# ══════════════════════════════════════════════════════════════════════
# Top Summary, Regional Speedometer Gauge & Branch Strip Distribution
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/summary", methods=["GET"])
@login_required
def api_rhi_summary():
    if request.args.get("refresh") == "1":
        clear_rhi_caches()
    yearweeks = _get_yearweeks()
    if not yearweeks:
        return jsonify({"status": "error", "message": "No data available in sow.rhi_score"}), 404

    try:
        req_yw = request.args.get("weeknum", type=int)
        target_yw = req_yw if req_yw and req_yw in yearweeks else yearweeks[0]
    except Exception:
        target_yw = yearweeks[0]

    cache_key = f"summary:{target_yw}"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)

    curr_idx = yearweeks.index(target_yw)
    prev_yw = yearweeks[curr_idx + 1] if curr_idx + 1 < len(yearweeks) else None

    branch_map = _get_branch_map()
    target_val, target_label = _get_target_for_week(target_yw)

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            weeks_to_query = [target_yw] + ([prev_yw] if prev_yw else [])

            # Total and Good cells count per week
            cur.execute(
                """
                SELECT weeknum,
                       COUNT(*) AS total_cells,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good_cells,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail_cells
                FROM sow.rhi_score
                WHERE weeknum = ANY(%s)
                GROUP BY weeknum
                """,
                (weeks_to_query,)
            )
            cell_counts = {r[0]: {"total": r[1], "good": r[2], "fail": r[3]} for r in cur.fetchall()}

            # Technology breakdown (2G, 4G, 5G)
            cur.execute(
                """
                SELECT tech,
                       COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail
                FROM sow.rhi_score
                WHERE weeknum = %s
                GROUP BY tech
                """,
                (target_yw,)
            )
            tech_raw = cur.fetchall()
            tech_summary = {}
            for t_code, tot, gd, fl in tech_raw:
                t_key = (t_code or "").strip().lower()
                tech_summary[t_key] = {
                    "total": tot or 0,
                    "good": gd or 0,
                    "fail": fl or 0,
                    "pct": round((gd * 100.0 / tot), 2) if tot else 0.0
                }

            # Branch / NOP counts via kabupaten mapping
            cur.execute(
                """
                SELECT UPPER(TRIM(kabupaten)),
                       COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail
                FROM sow.rhi_score
                WHERE weeknum = %s AND kabupaten IS NOT NULL
                GROUP BY UPPER(TRIM(kabupaten))
                """,
                (target_yw,)
            )
            kab_counts = cur.fetchall()

        curr_data = cell_counts.get(target_yw, {"total": 0, "good": 0, "fail": 0})
        prev_data = cell_counts.get(prev_yw, {"total": 0, "good": 0, "fail": 0}) if prev_yw else None

        good_cell_pct = round(curr_data["good"] * 100.0 / curr_data["total"], 2) if curr_data["total"] else 0.0
        prev_good_cell_pct = (
            round(prev_data["good"] * 100.0 / prev_data["total"], 2) if (prev_data and prev_data["total"]) else None
        )
        wow_delta = round(good_cell_pct - prev_good_cell_pct, 2) if prev_good_cell_pct is not None else None
        diff_target_pp = round(good_cell_pct - target_val, 2)

        needed_recover = 0
        if good_cell_pct < target_val and curr_data["total"]:
            needed_good = int((target_val / 100.0) * curr_data["total"]) + 1
            needed_recover = max(0, needed_good - curr_data["good"])

        prev_label = f"W{str(prev_yw)[4:]}" if prev_yw else "-"

        # Aggregate to NOP branches: AMBON, JAYAPURA, SORONG, TIMIKA
        branch_agg = defaultdict(lambda: {"total": 0, "good": 0, "fail": 0})
        for kab, tot, gd, fl in kab_counts:
            b_name = branch_map.get(kab, "OTHER")
            branch_agg[b_name]["total"] += tot
            branch_agg[b_name]["good"] += gd
            branch_agg[b_name]["fail"] += fl

        regional_gauge = []
        for b_name in ["AMBON", "JAYAPURA", "SORONG", "TIMIKA"]:
            stats = branch_agg.get(b_name, {"total": 0, "good": 0, "fail": 0})
            pct = round(stats["good"] * 100.0 / stats["total"], 2) if stats["total"] else 0.0
            regional_gauge.append({
                "name": b_name,
                "value": pct,
                "total": stats["total"],
                "good": stats["good"],
                "fail": stats["fail"],
                "is_above": pct >= target_val
            })

        # Build Executive Summary Highlights
        quarter_str = _quarter_for_week(target_yw)
        yw_w = str(target_yw)[4:]
        yw_y = str(target_yw)[:4]
        fail_cell_pct = round(100.0 - good_cell_pct, 2)

        b1 = f"In Week W{yw_w} ({yw_y}), total monitored cells across PUMA is {curr_data['total']:,} cells, consisting of {curr_data['good']:,} Good cells ({good_cell_pct:.2f}%) and {curr_data['fail']:,} Fail cells ({fail_cell_pct:.2f}%)."

        if good_cell_pct >= target_val:
            b2 = f"Target RHI for PUMA in {quarter_str} is {target_val:.2f}%, successfully passing the target with a margin of +{diff_target_pp:.2f} pp (0 cells required to recover)."
        else:
            b2 = f"Target RHI for PUMA in {quarter_str} is {target_val:.2f}%, breaching the target by {abs(diff_target_pp):.2f} pp. Need to recover {needed_recover:,} fail cells for optimization to reach target."

        if wow_delta is not None and prev_good_cell_pct is not None:
            sign = '+' if wow_delta > 0 else ''
            wow_desc = f"improved by {sign}{wow_delta:.2f} pp" if wow_delta > 0 else (f"degraded by {wow_delta:.2f} pp" if wow_delta < 0 else "stable at 0.00 pp")
            b3 = f"Week-over-Week (WoW) performance compared to {prev_label} has {wow_desc} (from {prev_good_cell_pct:.2f}% to {good_cell_pct:.2f}%)."
        else:
            b3 = f"Current overall performance stands at {good_cell_pct:.2f}% against network baseline."

        t4g = tech_summary.get('4g', {})
        t2g = tech_summary.get('2g', {})
        t5g = tech_summary.get('5g', {})
        b4 = f"Technology breakdown: 4G Network achieves {t4g.get('pct', 0):.2f}% ({t4g.get('good', 0):,} of {t4g.get('total', 0):,} cells), 2G Network achieves {t2g.get('pct', 0):.2f}% ({t2g.get('good', 0):,} of {t2g.get('total', 0):,} cells), and 5G Network achieves {t5g.get('pct', 0):.2f}%."

        if regional_gauge:
            best_branch = max(regional_gauge, key=lambda x: x["value"])
            lowest_branch = min(regional_gauge, key=lambda x: x["value"])
            passing_branches = [b["name"] for b in regional_gauge if b["value"] >= target_val]
            passing_str = ", ".join(passing_branches) if passing_branches else "None"
            b5 = f"Branch performance: TOP branch is {best_branch['name']} ({best_branch['value']:.2f}%), while lowest is {lowest_branch['name']} ({lowest_branch['value']:.2f}%). Branches passing target: {passing_str}."
        else:
            b5 = "Regional branch performance distribution is within active monitoring thresholds."

        bullets = [b1, b2, b3, b4, b5]

        res_data = {
            "status": "success",
            "yearweek": target_yw,
            "prev_yearweek": prev_yw,
            "yearweeks": yearweeks,
            "total_cells": curr_data["total"],
            "good_cells": curr_data["good"],
            "fail_cells": curr_data["fail"],
            "good_cell_pct": good_cell_pct,
            "prev_good_cell_pct": prev_good_cell_pct,
            "wow_delta": wow_delta,
            "target": target_val,
            "target_label": target_label,
            "baseline": 93.05,
            "quarter": _quarter_for_week(target_yw),
            "is_above_target": good_cell_pct >= target_val,
            "diff_target_pp": diff_target_pp,
            "diff_prev_pp": wow_delta,
            "needed_recover": needed_recover,
            "prev_label": prev_label,
            "tech_summary": tech_summary,
            "regional_gauge": regional_gauge,
            "bullets": bullets,
        }
        _RHI_CACHE[cache_key] = (res_data, now + _RHI_CACHE_TTL)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error calculating RHI summary: %s", e)
        return jsonify({"status": "error", "message": "Failed to calculate summary"}), 500


# ══════════════════════════════════════════════════════════════════════
# Indicator Cards: Tech & Key KPI Health with WoW & Sparklines
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/indicator-cards", methods=["GET"])
@rhi.route("/api/sow/rhi/indicator_cards", methods=["GET"])
@rhi.route("/api/okr/rhi/indicator-cards", methods=["GET"])
@rhi.route("/api/okr/rhi/indicator_cards", methods=["GET"])
@login_required
def api_rhi_indicator_cards():
    yearweeks = _get_yearweeks()
    if not yearweeks:
        return jsonify({"status": "error", "message": "No data"}), 404

    try:
        req_yw = request.args.get("weeknum", type=int)
        target_yw = req_yw if req_yw and req_yw in yearweeks else yearweeks[0]
    except Exception:
        target_yw = yearweeks[0]

    cache_key = f"indicator_cards:{target_yw}"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)

    curr_idx = yearweeks.index(target_yw)
    prev_yw = yearweeks[curr_idx + 1] if curr_idx + 1 < len(yearweeks) else None
    spark_yws = list(reversed(yearweeks[curr_idx:curr_idx + 8]))

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            # 1. Tech Health across sparkline weeks
            cur.execute(
                """
                SELECT weeknum, tech,
                       COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good
                FROM sow.rhi_score
                WHERE weeknum = ANY(%s)
                GROUP BY weeknum, tech
                ORDER BY weeknum ASC
                """,
                (spark_yws,)
            )
            tech_hist = defaultdict(lambda: defaultdict(lambda: {"total": 0, "good": 0}))
            for yw, tech, tot, gd in cur.fetchall():
                t_clean = (tech or "").strip().lower()
                tech_hist[t_clean][yw] = {"total": tot or 0, "good": gd or 0}

            # 2. Key Indicator Averages & Fail Cell counts for target and prev week
            weeks_compare = [target_yw] + ([prev_yw] if prev_yw else [])
            ind_exprs = []
            for ind in RHI_INDICATORS:
                ind_exprs.append(
                    f'ROUND(AVG("{ind["val_col"]}") FILTER (WHERE tech = \'{ind["tech"]}\')::numeric, 2) AS avg_{ind["id"]}'
                )
                ind_exprs.append(
                    f'COUNT(*) FILTER (WHERE tech = \'{ind["tech"]}\' AND "{ind["fail_col"]}" = 1) AS fail_{ind["id"]}'
                )
            q_ind = f"""
                SELECT weeknum, {", ".join(ind_exprs)}
                FROM sow.rhi_score
                WHERE weeknum = ANY(%s)
                GROUP BY weeknum
                ORDER BY weeknum ASC
            """
            cur.execute(q_ind, (spark_yws,))
            ind_by_week = {}
            for r in cur.fetchall():
                yw = r[0]
                idx = 1
                row_map = {}
                for ind in RHI_INDICATORS:
                    row_map[f'avg_{ind["id"]}'] = float(r[idx]) if r[idx] is not None else 0.0
                    row_map[f'fail_{ind["id"]}'] = int(r[idx + 1]) if r[idx + 1] is not None else 0
                    idx += 2
                ind_by_week[yw] = row_map

            # 3. Top Fail Root Cause in target week for 2G, 4G, 5G
            cur.execute(
                """
                SELECT tech,
                       COALESCE(most_fail_kpi_4g, most_fail_kpi_2g, most_fail_kpi_5g, 'Other') AS reason,
                       COUNT(*) AS cnt
                FROM sow.rhi_score
                WHERE weeknum = %s AND UPPER(TRIM(remark_week)) = 'FAIL'
                GROUP BY tech, reason
                ORDER BY tech, cnt DESC
                """,
                (target_yw,)
            )
            top_reasons = {}
            for tech, r_code, cnt in cur.fetchall():
                t_clean = (tech or "").strip().lower()
                if t_clean not in top_reasons:
                    top_reasons[t_clean] = {
                        "code": r_code,
                        "label": FAIL_REASON_LABELS.get(r_code, r_code or "None"),
                        "count": cnt
                    }

        # Format Tech Cards
        tech_cards = []
        for t_code, t_name, t_color in [("2g", "2G Network", "#db2777"), ("4g", "4G Network", "#2563eb"), ("5g", "5G Network", "#10b981")]:
            t_curr = tech_hist[t_code].get(target_yw, {"total": 0, "good": 0})
            t_prev = tech_hist[t_code].get(prev_yw, {"total": 0, "good": 0}) if prev_yw else None

            pct_curr = round(t_curr["good"] * 100.0 / t_curr["total"], 2) if t_curr["total"] else 0.0
            pct_prev = round(t_prev["good"] * 100.0 / t_prev["total"], 2) if (t_prev and t_prev["total"]) else None
            delta = round(pct_curr - pct_prev, 2) if pct_prev is not None else None

            spark = []
            for syw in spark_yws:
                s_stat = tech_hist[t_code].get(syw, {"total": 0, "good": 0})
                s_pct = round(s_stat["good"] * 100.0 / s_stat["total"], 2) if s_stat["total"] else 0.0
                spark.append(s_pct)

            top_r = top_reasons.get(t_code, {"code": "None", "label": "None", "count": 0})
            top_cnt = top_r.get("count", 0)
            tot_cells = t_curr["total"]
            fail_cells = tot_cells - t_curr["good"]
            top_fail_pct = round(top_cnt * 100.0 / tot_cells, 2) if tot_cells else 0.0
            top_fail_share = round(top_cnt * 100.0 / fail_cells, 2) if fail_cells else 0.0

            tech_cards.append({
                "tech": t_code,
                "title": t_name,
                "color": t_color,
                "total": t_curr["total"],
                "good": t_curr["good"],
                "fail": fail_cells,
                "pct": pct_curr,
                "prev_pct": pct_prev,
                "delta": delta,
                "sparkline": spark,
                "top_fail_label": top_r["label"],
                "top_fail_count": top_cnt,
                "top_fail_pct": top_fail_pct,
                "top_fail_share": top_fail_share,
            })

        # Format Indicator Cards
        curr_ind = ind_by_week.get(target_yw, {})
        prev_ind = ind_by_week.get(prev_yw, {}) if prev_yw else {}
        
        tech_fail_map = {c["tech"].lower(): c["fail"] for c in tech_cards}
        total_fail_4g = tech_fail_map.get("4g", 0)
        total_fail_2g = tech_fail_map.get("2g", 0)
        total_fail_all = sum(c["fail"] for c in tech_cards)

        indicator_cards = []
        for ind in RHI_INDICATORS:
            val_c = curr_ind.get(f'avg_{ind["id"]}', 0.0)
            val_p = prev_ind.get(f'avg_{ind["id"]}', None) if prev_yw else None
            fail_c = curr_ind.get(f'fail_{ind["id"]}', 0)
            fail_p = prev_ind.get(f'fail_{ind["id"]}', None) if prev_yw else None
            delta = round(val_c - val_p, 2) if val_p is not None else None
            delta_fail = (fail_c - fail_p) if fail_p is not None else None
            
            tech_key = ind["tech"].lower()
            tech_fails = tech_fail_map.get(tech_key, 0)
            fail_share = round((fail_c * 100.0 / tech_fails), 1) if tech_fails else 0.0

            # 6-8 week trend of fail cells and scores
            fail_trend = [ind_by_week.get(syw, {}).get(f'fail_{ind["id"]}', 0) for syw in spark_yws]
            score_trend = [ind_by_week.get(syw, {}).get(f'avg_{ind["id"]}', 0.0) for syw in spark_yws]

            indicator_cards.append({
                "id": ind["id"],
                "title": ind["title"],
                "tech": ind["tech"].upper(),
                "color": ind["color"],
                "value": val_c,
                "unit": ind["unit"],
                "delta": delta,
                "fail_count": fail_c,
                "delta_fail": delta_fail,
                "tech_fail_total": tech_fails,
                "fail_share": fail_share,
                "fail_trend": fail_trend,
                "score_trend": score_trend,
                "spark_weeks": [f"W{str(w)[4:]}" for w in spark_yws],
            })

        res_data = {
            "status": "success",
            "yearweek": target_yw,
            "prev_yearweek": prev_yw,
            "tech_cards": tech_cards,
            "indicator_cards": indicator_cards,
            "total_fail_4g": total_fail_4g,
            "total_fail_2g": total_fail_2g,
            "total_fail_all": total_fail_all,
        }
        _RHI_CACHE[cache_key] = (res_data, now + _RHI_CACHE_TTL)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error loading RHI indicator cards: %s", e)
        return jsonify({"status": "error", "message": "Failed to load indicator cards"}), 500


# ══════════════════════════════════════════════════════════════════════
# Dual-Axis Macro Trend Achievement Chart
# Stacked Bars: Good & Fail counts; Line: % Good Cell + Target
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/achievement", methods=["GET"])
@login_required
def api_rhi_achievement():
    yearweeks = _get_yearweeks()
    if not yearweeks:
        return jsonify({"status": "error", "message": "No data"}), 404

    try:
        start_yw = request.args.get("start_yw", type=int) or (yearweeks[-1] if yearweeks else 202637)
        end_yw = request.args.get("end_yw", type=int) or yearweeks[0]
    except Exception:
        start_yw, end_yw = (yearweeks[-1] if yearweeks else 202637), yearweeks[0]

    cache_key = f"achievement:{start_yw}:{end_yw}"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(
                """
                SELECT weeknum,
                       COUNT(*) AS total_cells,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good_cells,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail_cells,
                       COUNT(*) FILTER (WHERE tech = '2g' AND UPPER(TRIM(remark_week)) = 'FAIL') AS fail_2g,
                       COUNT(*) FILTER (WHERE tech = '4g' AND UPPER(TRIM(remark_week)) = 'FAIL') AS fail_4g,
                       COUNT(*) FILTER (WHERE tech = '5g' AND UPPER(TRIM(remark_week)) = 'FAIL') AS fail_5g
                FROM sow.rhi_score
                WHERE weeknum >= %s AND weeknum <= %s
                GROUP BY weeknum
                ORDER BY weeknum ASC
                """,
                (start_yw, end_yw)
            )
            rows = cur.fetchall()

        weeks, labels, good_counts, fail_counts, total_counts, good_pct = [], [], [], [], [], []
        fail_2g_series, fail_4g_series, fail_5g_series = [], [], []
        for yw, total, good, fail, f2g, f4g, f5g in rows:
            weeks.append(yw)
            yw_s = str(yw)
            labels.append(f"W{yw_s[4:]}")
            good_counts.append(good or 0)
            fail_counts.append(fail or 0)
            total_counts.append(total or 0)
            good_pct.append(round(good * 100.0 / total, 2) if total else 0.0)
            fail_2g_series.append(f2g or 0)
            fail_4g_series.append(f4g or 0)
            fail_5g_series.append(f5g or 0)

        all_targets = _get_all_targets()
        target_series = []
        for yw in weeks:
            t = all_targets.get(int(str(yw)[:4]), {}).get(_quarter_for_week(yw))
            target_series.append(t if t is not None else 93.55)

        target_val, _ = _get_target_for_week(end_yw if weeks else (yearweeks[0]))

        res_data = {
            "status": "success",
            "weeks": weeks,
            "labels": labels,
            "good_counts": good_counts,
            "fail_counts": fail_counts,
            "total_counts": total_counts,
            "good_pct_series": good_pct,
            "target_series": target_series,
            "target": target_val,
            "fail_2g": fail_2g_series,
            "fail_4g": fail_4g_series,
            "fail_5g": fail_5g_series,
        }
        _RHI_CACHE[cache_key] = (res_data, now + _RHI_CACHE_TTL)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error calculating RHI achievement: %s", e)
        return jsonify({"status": "error", "message": "Failed to compute achievement"}), 500


# ══════════════════════════════════════════════════════════════════════
# Multi-Tech Trend & Top Fail Root Cause Breakdown
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/trend", methods=["GET"])
@login_required
def api_rhi_trend():
    yearweeks = _get_yearweeks()
    if not yearweeks:
        return jsonify({"status": "error", "message": "No data"}), 404

    try:
        start_yw = request.args.get("start_yw", type=int) or (yearweeks[-1] if yearweeks else 202637)
        end_yw = request.args.get("end_yw", type=int) or yearweeks[0]
    except Exception:
        start_yw, end_yw = (yearweeks[-1] if yearweeks else 202637), yearweeks[0]

    cache_key = f"trend:{start_yw}:{end_yw}"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            # 1. Tech series per week
            cur.execute(
                """
                SELECT weeknum, tech,
                       COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good
                FROM sow.rhi_score
                WHERE weeknum >= %s AND weeknum <= %s
                GROUP BY weeknum, tech
                ORDER BY weeknum ASC
                """,
                (start_yw, end_yw)
            )
            raw_tech = cur.fetchall()

            # 2. Top fail reasons per week
            cur.execute(
                """
                SELECT weeknum,
                       COALESCE(most_fail_kpi_4g, most_fail_kpi_2g, most_fail_kpi_5g, 'Other') AS reason,
                       COUNT(*) AS cnt
                FROM sow.rhi_score
                WHERE weeknum >= %s AND weeknum <= %s AND UPPER(TRIM(remark_week)) = 'FAIL'
                GROUP BY weeknum, reason
                ORDER BY weeknum ASC, cnt DESC
                """,
                (start_yw, end_yw)
            )
            raw_reasons = cur.fetchall()

        # Build chronological weeks list
        weeks_set = sorted(list({r[0] for r in raw_tech}))
        labels = [f"W{str(w)[4:]}" for w in weeks_set]

        tech_data = {
            "overall": [],
            "2g": [],
            "4g": [],
            "5g": []
        }
        tech_map = defaultdict(lambda: defaultdict(lambda: {"total": 0, "good": 0}))
        weekly_overall = defaultdict(lambda: {"total": 0, "good": 0})
        for yw, tech, tot, gd in raw_tech:
            t_clean = (tech or "").strip().lower()
            tech_map[t_clean][yw] = {"total": tot, "good": gd}
            weekly_overall[yw]["total"] += tot
            weekly_overall[yw]["good"] += gd

        for yw in weeks_set:
            o = weekly_overall[yw]
            tech_data["overall"].append(round(o["good"] * 100.0 / o["total"], 2) if o["total"] else 0.0)
            for tc in ["2g", "4g", "5g"]:
                t = tech_map[tc][yw]
                tech_data[tc].append(round(t["good"] * 100.0 / t["total"], 2) if t["total"] else 0.0)

        # Fail reasons matrix
        # Get top 6 reasons overall
        reasons_totals = defaultdict(int)
        for yw, reason, cnt in raw_reasons:
            reasons_totals[reason] += cnt
        top_reasons_sorted = sorted(reasons_totals.keys(), key=lambda r: reasons_totals[r], reverse=True)[:6]

        reason_by_week = defaultdict(lambda: defaultdict(int))
        for yw, reason, cnt in raw_reasons:
            reason_by_week[reason][yw] = cnt

        reason_series = []
        for r_code in top_reasons_sorted:
            r_label = FAIL_REASON_LABELS.get(r_code, r_code)
            series_data = [reason_by_week[r_code][yw] for yw in weeks_set]
            reason_series.append({
                "code": r_code,
                "label": r_label,
                "data": series_data,
            })

        res_data = {
            "status": "success",
            "weeks": weeks_set,
            "labels": labels,
            "tech_data": tech_data,
            "reason_series": reason_series,
        }
        _RHI_CACHE[cache_key] = (res_data, now + _RHI_CACHE_TTL)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error loading RHI trend: %s", e)
        return jsonify({"status": "error", "message": "Failed to compute trend"}), 500


# ══════════════════════════════════════════════════════════════════════
# Geo Options & Geo Trend (Kabupaten & Kecamatan Filter)
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/geo-options", methods=["GET"])
@login_required
def api_rhi_geo_options():
    cache_key = "geo_options"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute("""
                SELECT TRIM(kabupaten), TRIM(kecamatan) FROM sow.rhi_score
                WHERE weeknum = (SELECT MAX(weeknum) FROM sow.rhi_score)
                  AND kabupaten IS NOT NULL AND kecamatan IS NOT NULL
                GROUP BY TRIM(kabupaten), TRIM(kecamatan) ORDER BY 1, 2
            """)
            kec_map = {}
            for kab, kec in cur.fetchall():
                if kab:
                    kec_map.setdefault(kab, []).append(kec)
            kabs = sorted(list(kec_map.keys()))
        res_data = {"status": "success", "kabupaten": kabs, "kecamatan_map": kec_map}
        _RHI_CACHE[cache_key] = (res_data, now + 86400)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error loading RHI geo options: %s", e)
        return jsonify({"status": "error", "message": "Failed to load geo options"}), 500


@rhi.route("/api/sow/rhi/geo-trend", methods=["GET"])
@login_required
def api_rhi_geo_trend():
    kabupaten = (request.args.get("kabupaten") or "").strip()
    kecamatan = (request.args.get("kecamatan") or "").strip()
    tech = (request.args.get("tech") or "all").strip().lower()

    where_parts, params = [], []
    if kabupaten:
        where_parts.append("UPPER(TRIM(kabupaten)) = UPPER(TRIM(%s))")
        params.append(kabupaten)
    if kecamatan:
        where_parts.append("UPPER(TRIM(kecamatan)) = UPPER(TRIM(%s))")
        params.append(kecamatan)
    if tech in ["2g", "4g", "5g"]:
        where_parts.append("LOWER(TRIM(tech)) = %s")
        params.append(tech)

    geo_where = (" AND ".join(where_parts))
    geo_where = f"AND {geo_where}" if geo_where else ""

    cache_key = f"geo_trend:{kabupaten}:{kecamatan}:{tech}"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)

    try:
        with db_query(get_postgres_connection) as (conn, cur):
            q_pct = f"""
                SELECT weeknum, COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail
                FROM sow.rhi_score WHERE 1=1 {geo_where}
                GROUP BY weeknum ORDER BY weeknum ASC
            """
            cur.execute(q_pct, tuple(params))
            pct_rows = cur.fetchall()

            # Latest week summary with geo filter
            q_sum = f"""
                SELECT COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail,
                       COUNT(*) FILTER (WHERE tech = '2g') AS total_2g,
                       COUNT(*) FILTER (WHERE tech = '2g' AND UPPER(TRIM(remark_week)) = 'GOOD') AS good_2g,
                       COUNT(*) FILTER (WHERE tech = '4g') AS total_4g,
                       COUNT(*) FILTER (WHERE tech = '4g' AND UPPER(TRIM(remark_week)) = 'GOOD') AS good_4g,
                       COUNT(*) FILTER (WHERE tech = '5g') AS total_5g,
                       COUNT(*) FILTER (WHERE tech = '5g' AND UPPER(TRIM(remark_week)) = 'GOOD') AS good_5g
                FROM sow.rhi_score
                WHERE weeknum = (SELECT MAX(weeknum) FROM sow.rhi_score WHERE 1=1 {geo_where})
                {geo_where}
            """
            cur.execute(q_sum, tuple(params + params))
            srow = cur.fetchone()

        weeks, labels, good_pct = [], [], []
        for yw, total, good, fail in pct_rows:
            weeks.append(yw)
            yw_s = str(yw)
            labels.append(f"W{yw_s[4:]}")
            good_pct.append(round(good * 100.0 / total, 2) if total else 0.0)

        all_targets = _get_all_targets()
        target_series = []
        for yw in weeks:
            t = all_targets.get(int(str(yw)[:4]), {}).get(_quarter_for_week(yw))
            target_series.append(t if t is not None else 93.55)

        latest_yw = weeks[-1] if weeks else None
        tgt_latest, _ = _get_target_for_week(latest_yw) if latest_yw else (93.55, "")
        tot = srow[0] or 0 if srow else 0
        gd = srow[1] or 0 if srow else 0
        fl = srow[2] or 0 if srow else 0
        t2g, g2g = (srow[3] or 0, srow[4] or 0) if srow else (0, 0)
        t4g, g4g = (srow[5] or 0, srow[6] or 0) if srow else (0, 0)
        t5g, g5g = (srow[7] or 0, srow[8] or 0) if srow else (0, 0)

        pct_latest = round(gd * 100.0 / tot, 2) if tot else 0.0
        pct_2g = round(g2g * 100.0 / t2g, 2) if t2g else 0.0
        pct_4g = round(g4g * 100.0 / t4g, 2) if t4g else 0.0
        pct_5g = round(g5g * 100.0 / t5g, 2) if t5g else 0.0
        diff_tgt = round(pct_latest - tgt_latest, 2)

        res_data = {
            "status": "success",
            "weeks": weeks,
            "labels": labels,
            "good_pct_series": good_pct,
            "target_series": target_series,
            "summary": {
                "weeknum": latest_yw,
                "total": tot,
                "good": gd,
                "fail": fl,
                "pct": pct_latest,
                "target": tgt_latest,
                "diff_target": diff_tgt,
                "is_above": pct_latest >= tgt_latest,
                "kabupaten": kabupaten or "All",
                "kecamatan": kecamatan or "All",
                "tech": tech.upper(),
                "total_2g": t2g, "good_2g": g2g, "pct_2g": pct_2g,
                "total_4g": t4g, "good_4g": g4g, "pct_4g": pct_4g,
                "total_5g": t5g, "good_5g": g5g, "pct_5g": pct_5g,
            }
        }
        _RHI_CACHE[cache_key] = (res_data, now + _RHI_CACHE_TTL)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error calculating RHI geo trend: %s", e)
        return jsonify({"status": "error", "message": "Failed to compute geo trend"}), 500


# ══════════════════════════════════════════════════════════════════════
# Comparison: NOP Level WoW + Technology/Root Cause WoW Table
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/comparison", methods=["GET"])
@login_required
def api_rhi_comparison():
    yearweeks = _get_yearweeks()
    if len(yearweeks) < 1:
        return jsonify({"status": "error", "message": "No data"}), 404

    try:
        req_yw = request.args.get("weeknum", type=int)
        curr_yw = req_yw if req_yw and req_yw in yearweeks else yearweeks[0]
    except Exception:
        curr_yw = yearweeks[0]

    curr_idx = yearweeks.index(curr_yw) if curr_yw in yearweeks else 0
    prev_yw = yearweeks[curr_idx + 1] if curr_idx + 1 < len(yearweeks) else None
    if prev_yw is None:
        return jsonify({"status": "error", "message": "Need at least 2 weeks for WoW comparison"}), 400

    target_val, _ = _get_target_for_week(curr_yw)
    cache_key = f"comparison:{curr_yw}:{prev_yw}"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)

    try:
        branch_map = _get_branch_map()
        with db_query(get_postgres_connection) as (conn, cur):
            # 1. Kabupaten counts
            cur.execute("""
                SELECT UPPER(TRIM(kabupaten)), weeknum,
                       COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail
                FROM sow.rhi_score
                WHERE weeknum = ANY(%s) AND kabupaten IS NOT NULL
                GROUP BY UPPER(TRIM(kabupaten)), weeknum
            """, ([curr_yw, prev_yw],))
            kab_rows = cur.fetchall()

            # 2. Tech WoW
            cur.execute("""
                SELECT tech, weeknum,
                       COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail
                FROM sow.rhi_score
                WHERE weeknum = ANY(%s)
                GROUP BY tech, weeknum
            """, ([curr_yw, prev_yw],))
            tech_rows = cur.fetchall()

            # 3. Top Fail Causes WoW
            cur.execute("""
                SELECT COALESCE(most_fail_kpi_4g, most_fail_kpi_2g, most_fail_kpi_5g, 'Other') AS reason,
                       weeknum,
                       COUNT(*) AS cnt
                FROM sow.rhi_score
                WHERE weeknum = ANY(%s) AND UPPER(TRIM(remark_week)) = 'FAIL'
                GROUP BY 1, 2
            """, ([curr_yw, prev_yw],))
            reason_rows = cur.fetchall()

        # Aggregate kabupaten -> branch
        branches = ["AMBON", "JAYAPURA", "SORONG", "TIMIKA"]
        agg = {b: {curr_yw: [0, 0, 0], prev_yw: [0, 0, 0]} for b in branches}
        puma = {curr_yw: [0, 0, 0], prev_yw: [0, 0, 0]}

        for kab, yw, total, good, fail in kab_rows:
            b = branch_map.get(kab, "OTHER")
            puma[yw][0] += total or 0
            puma[yw][1] += good or 0
            puma[yw][2] += fail or 0
            if b in agg:
                agg[b][yw][0] += total or 0
                agg[b][yw][1] += good or 0
                agg[b][yw][2] += fail or 0

        def _row(nop, d_curr, d_prev):
            tot_c, gd_c, fl_c = d_curr
            tot_p, gd_p, fl_p = d_prev
            ach_c = round(gd_c * 100.0 / tot_c, 2) if tot_c else 0.0
            ach_p = round(gd_p * 100.0 / tot_p, 2) if tot_p else 0.0
            d_tgt = round(ach_c - target_val, 2)
            wow = round(ach_c - ach_p, 2)
            return {
                "nop": nop,
                "target": target_val,
                "prev_total": tot_p,
                "prev_good": gd_p,
                "prev_bad": fl_p,
                "prev_ach": ach_p,
                "curr_total": tot_c,
                "curr_good": gd_c,
                "curr_bad": fl_c,
                "curr_ach": ach_c,
                "delta_target": d_tgt,
                "ach_wow": wow,
                "st_target": "Pass" if ach_c >= target_val else "Not Pass",
                "st_wow": "Improved" if wow > 0 else ("Stable" if wow == 0 else "Degraded"),
            }

        rows = [_row(b, agg[b][curr_yw], agg[b][prev_yw]) for b in branches]
        rows.append(_row("PUMA", puma[curr_yw], puma[prev_yw]))

        # Tech WoW table
        tech_agg = defaultdict(lambda: {curr_yw: [0, 0, 0], prev_yw: [0, 0, 0]})
        for tech, yw, tot, gd, fl in tech_rows:
            t_clean = (tech or "").strip().lower()
            tech_agg[t_clean][yw] = [tot or 0, gd or 0, fl or 0]

        tech_wow = []
        for t_code in ["2g", "4g", "5g"]:
            d_c = tech_agg[t_code][curr_yw]
            d_p = tech_agg[t_code][prev_yw]
            ach_c = round(d_c[1] * 100.0 / d_c[0], 2) if d_c[0] else 0.0
            ach_p = round(d_p[1] * 100.0 / d_p[0], 2) if d_p[0] else 0.0
            delta = round(ach_c - ach_p, 2)
            tech_wow.append({
                "tech": t_code.upper(),
                "prev_total": d_p[0],
                "prev_good": d_p[1],
                "prev_ach": ach_p,
                "curr_total": d_c[0],
                "curr_good": d_c[1],
                "curr_ach": ach_c,
                "delta": delta,
                "remark": "Improve" if delta > 0 else ("Stable" if delta == 0 else "Degrade")
            })

        # Top Fail Root Causes WoW table
        r_agg = defaultdict(lambda: {curr_yw: 0, prev_yw: 0})
        for r_code, yw, cnt in reason_rows:
            r_agg[r_code][yw] = cnt or 0

        # Sort by current week count
        sorted_reasons = sorted(r_agg.keys(), key=lambda k: r_agg[k][curr_yw], reverse=True)[:8]
        reasons_wow = []
        for r_code in sorted_reasons:
            before = r_agg[r_code][prev_yw]
            after = r_agg[r_code][curr_yw]
            delta = after - before
            delta_pct = round(delta * 100.0 / before, 2) if before else (0.0 if after == 0 else 100.0)
            reasons_wow.append({
                "code": r_code,
                "title": FAIL_REASON_LABELS.get(r_code, r_code),
                "before": before,
                "after": after,
                "delta": delta,
                "delta_pct": delta_pct,
                # For failures, decrease is improvement
                "remark": "Improve" if delta < 0 else ("Stable" if delta == 0 else "Degrade"),
            })

        res_data = {
            "status": "success",
            "curr_week": curr_yw,
            "prev_week": prev_yw,
            "target": target_val,
            "rows": rows,
            "tech_wow": tech_wow,
            "reasons_wow": reasons_wow,
        }
        _RHI_CACHE[cache_key] = (res_data, now + _RHI_CACHE_TTL)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error calculating RHI comparison: %s", e)
        return jsonify({"status": "error", "message": "Failed to compute comparison"}), 500


# ══════════════════════════════════════════════════════════════════════
# Kabupaten Level WoW Comparison Table
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/kabupaten-comparison", methods=["GET"])
@login_required
def api_rhi_kabupaten_comparison():
    yearweeks = _get_yearweeks()
    if len(yearweeks) < 1:
        return jsonify({"status": "error", "message": "No data"}), 404

    try:
        req_yw = request.args.get("weeknum", type=int)
        curr_yw = req_yw if req_yw and req_yw in yearweeks else yearweeks[0]
    except Exception:
        curr_yw = yearweeks[0]

    curr_idx = yearweeks.index(curr_yw) if curr_yw in yearweeks else 0
    prev_yw = yearweeks[curr_idx + 1] if curr_idx + 1 < len(yearweeks) else None
    if prev_yw is None:
        return jsonify({"status": "error", "message": "Need at least 2 weeks for WoW comparison"}), 400

    target_val, _ = _get_target_for_week(curr_yw)
    branch_filter = (request.args.get("branch") or "ALL").strip().upper()
    tech_filter = (request.args.get("tech") or "ALL").strip().lower()

    cache_key = f"kab_comp:{curr_yw}:{prev_yw}:{branch_filter}:{tech_filter}"
    now = time.time()
    if cache_key in _RHI_CACHE:
        val, exp = _RHI_CACHE[cache_key]
        if now < exp:
            return jsonify(val)

    tech_clause = ""
    query_params = [[curr_yw, prev_yw]]
    if tech_filter in ["2g", "4g", "5g"]:
        tech_clause = "AND LOWER(TRIM(tech)) = %s"
        query_params.append(tech_filter)

    try:
        branch_map = _get_branch_map()
        with db_query(get_postgres_connection) as (conn, cur):
            cur.execute(f"""
                SELECT UPPER(TRIM(kabupaten)), weeknum,
                       COUNT(*) AS total,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'GOOD') AS good,
                       COUNT(*) FILTER (WHERE UPPER(TRIM(remark_week)) = 'FAIL') AS fail
                FROM sow.rhi_score
                WHERE weeknum = ANY(%s) AND kabupaten IS NOT NULL {tech_clause}
                GROUP BY UPPER(TRIM(kabupaten)), weeknum
                ORDER BY 1, 2
            """, tuple(query_params))
            rows = cur.fetchall()

        kab_data = defaultdict(lambda: {curr_yw: [0, 0, 0], prev_yw: [0, 0, 0]})
        for kab, yw, total, good, fail in rows:
            kab_data[kab][yw] = [total or 0, good or 0, fail or 0]

        result_rows = []
        for kab, w_dict in sorted(kab_data.items()):
            b_name = branch_map.get(kab, "OTHER")
            if branch_filter != "ALL" and b_name != branch_filter:
                continue

            tot_c, gd_c, fl_c = w_dict[curr_yw]
            tot_p, gd_p, fl_p = w_dict[prev_yw]

            ach_c = round(gd_c * 100.0 / tot_c, 2) if tot_c else 0.0
            ach_p = round(gd_p * 100.0 / tot_p, 2) if tot_p else 0.0
            d_tgt = round(ach_c - target_val, 2)
            wow = round(ach_c - ach_p, 2)

            result_rows.append({
                "branch": b_name,
                "kabupaten": kab,
                "target": target_val,
                "prev_total": tot_p,
                "prev_good": gd_p,
                "prev_bad": fl_p,
                "prev_ach": ach_p,
                "curr_total": tot_c,
                "curr_good": gd_c,
                "curr_bad": fl_c,
                "curr_ach": ach_c,
                "delta_target": d_tgt,
                "ach_wow": wow,
                "st_target": "Pass" if ach_c >= target_val else "Not Pass",
                "st_wow": "Improved" if wow > 0 else ("Stable" if wow == 0 else "Degraded"),
            })

        # Sort primarily by Branch name, then by Kabupaten
        result_rows.sort(key=lambda r: (r["branch"], r["kabupaten"]))

        res_data = {
            "status": "success",
            "curr_week": curr_yw,
            "prev_week": prev_yw,
            "target": target_val,
            "rows": result_rows,
        }
        _RHI_CACHE[cache_key] = (res_data, now + _RHI_CACHE_TTL)
        return jsonify(res_data)
    except Exception as e:
        logger.exception("Error calculating RHI kabupaten comparison: %s", e)
        return jsonify({"status": "error", "message": "Failed to compute kabupaten comparison"}), 500


# ══════════════════════════════════════════════════════════════════════
# CSV Export
# ══════════════════════════════════════════════════════════════════════
@rhi.route("/api/sow/rhi/export", methods=["GET"])
@login_required
def api_rhi_export():
    yearweeks = _get_yearweeks()
    if not yearweeks:
        return jsonify({"status": "error", "message": "No data"}), 404

    start_yw = request.args.get("start_yw", type=int) or yearweeks[-1]
    end_yw = request.args.get("end_yw", type=int) or yearweeks[0]

    with db_query(get_postgres_connection) as (conn, cur):
        cur.execute("""
            SELECT weeknum, tech, regional, vendor, cell_name, site_id, kabupaten, kecamatan,
                   remark_week, total_fail_daily, cssr_tch, sdsr, ccsr_tch,
                   cssr_data_service, ccsr_data_service, volte_accessibility_sr, volte_retain_rate
            FROM sow.rhi_score
            WHERE weeknum >= %s AND weeknum <= %s
            ORDER BY weeknum DESC, kabupaten, cell_name
            LIMIT 50000
        """, (start_yw, end_yw))
        rows = cur.fetchall()

    headers = [
        "Weeknum", "Tech", "Regional", "Vendor", "Cell Name", "Site ID", "Kabupaten", "Kecamatan",
        "Remark Week", "Total Fail Daily", "CSSR TCH", "SDSR", "CCSR TCH",
        "CSSR Data Service", "CCSR Data Service", "VoLTE Acc SR", "VoLTE Retain Rate"
    ]
    filename = f"rhi_export_{start_yw}_{end_yw}.csv"
    return csv_response(rows, headers, filename)


# ══════════════════════════════════════════════════════════════════════
# Sample Template & Data Import for sow.rhi_score
# ══════════════════════════════════════════════════════════════════════
def _get_rhi_sample_rows():
    """Fetch sample rows from sow.rhi_score, or return synthetic fallback."""
    try:
        with db_query(get_postgres_connection) as (conn, cur):
            col_list = ", ".join([f'"{c}"' for c in RHI_COLUMNS])
            cur.execute(f"SELECT {col_list} FROM sow.rhi_score LIMIT 5")
            rows = cur.fetchall()
            if rows:
                return [dict(zip(RHI_COLUMNS, r)) for r in rows]
    except Exception as e:
        logger.warning("Could not fetch real sample rows for RHI: %s", e)

    return [{
        "weeknum": 202639,
        "tech": "4g",
        "regional": "PUMA",
        "vendor": "ZTE",
        "cell_name": "SAMPLE_CELL_01",
        "site_id": "SMP001",
        "kabupaten": "KOTA AMBON",
        "kecamatan": "SIRIMAU",
        "remark_week": "GOOD",
        "total_fail_daily": 0.0,
        "cssr_data_service": 99.8,
        "ccsr_data_service": 99.9,
    }]


@rhi.route("/api/sow/rhi/sample", methods=["GET"])
@rhi.route("/api/sow/rhi/sample-template", methods=["GET"])
@login_required
def api_sow_rhi_sample_template():
    """Download sample CSV or XLSX template for importing into sow.rhi_score."""
    file_format = (request.args.get("format") or "xlsx").lower()
    sample_rows = _get_rhi_sample_rows()

    if file_format == "xlsx":
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "rhi_score"

            ws.append(RHI_COLUMNS)
            for row_dict in sample_rows:
                ws.append([row_dict.get(c, "") for c in RHI_COLUMNS])

            out = io.BytesIO()
            wb.save(out)
            out.seek(0)
            return send_file(
                out,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                as_attachment=True,
                download_name="okr_rhi_sample_template.xlsx",
            )
        except ImportError:
            file_format = "csv"

    # CSV fallback
    si = io.StringIO()
    cw = csv.writer(si)
    cw.writerow(RHI_COLUMNS)
    for row_dict in sample_rows:
        cw.writerow([row_dict.get(c, "") for c in RHI_COLUMNS])

    output = make_response(si.getvalue())
    output.headers["Content-Disposition"] = "attachment; filename=okr_rhi_sample_template.csv"
    output.headers["Content-type"] = "text/csv; charset=utf-8"
    return output


class _RhiCountingReader:
    def __init__(self, fh, total_bytes, job_id):
        self._fh = fh
        self._total = max(1, int(total_bytes or 0))
        self._job_id = job_id

    def read(self, size=-1):
        chunk = self._fh.read(size)
        self._update(len(chunk))
        return chunk

    def _update(self, n):
        if n <= 0:
            return
        with _IMPORT_JOBS_LOCK:
            job = _IMPORT_JOBS.get(self._job_id)
            if not job:
                return
            job["_read"] = job.get("_read", 0) + n
            read_bytes = job["_read"]
        pct = 25 + int(70 * min(1.0, read_bytes / float(self._total)))
        _rhi_set_job(self._job_id, status="processing", pct=pct,
                     message=f"Uploading to database... {read_bytes/1048576:.1f}/{self._total/1048576:.1f} MB",
                     stage="copy")

    def readable(self):
        return True

    def writable(self):
        return False


def _rhi_set_job(job_id, **kw):
    with _IMPORT_JOBS_LOCK:
        job = _IMPORT_JOBS.get(job_id)
        if job:
            job.update(kw)


@rhi.route("/api/sow/rhi/import", methods=["POST"])
@csrf.exempt
@login_required
def api_sow_rhi_import():
    """Async import handler for RHI data."""
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file uploaded"}), 400

    file = request.files["file"]
    filename = (file.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx", ".xls")):
        return jsonify({"status": "error", "message": "Only CSV and Excel (.xlsx, .xls) files are supported"}), 400

    job_id = str(uuid4())
    fd, path = tempfile.mkstemp(suffix=os.path.splitext(filename)[1])
    with os.fdopen(fd, 'wb') as f:
        file.save(f)

    with _IMPORT_JOBS_LOCK:
        _IMPORT_JOBS[job_id] = {"status": "queued", "pct": 0, "message": "Queued...", "_read": 0}

    threading.Thread(target=_bg_rhi_import, args=(job_id, path, filename), daemon=True).start()
    return jsonify({"status": "queued", "job_id": job_id})


@rhi.route("/api/sow/rhi/import/progress/<job_id>")
@login_required
def api_sow_rhi_import_progress(job_id):
    with _IMPORT_JOBS_LOCK:
        job = _IMPORT_JOBS.get(job_id)
        if job is None:
            return jsonify({"status": "error", "message": "Job not found"}), 404
        snapshot = {k: v for k, v in job.items() if not k.startswith("_")}
    return jsonify(snapshot)


def _bg_rhi_import(job_id, file_path, filename):
    try:
        _rhi_set_job(job_id, status="processing", pct=5, message="Parsing file...", stage="parse")
        raw_rows = []
        is_csv = filename.lower().endswith(".csv")

        if is_csv:
            with open(file_path, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
                reader = csv.reader(fh)
                try:
                    headers = next(reader)
                except StopIteration:
                    _rhi_set_job(job_id, status="error", message="Uploaded CSV file is empty")
                    return
                idx = {str(h).strip().lower(): i for i, h in enumerate(headers)}
                for r in reader:
                    if any(r):
                        raw_rows.append([r[idx.get(c, -1)] if idx.get(c, -1) != -1 else "" for c in RHI_COLUMNS])
        else:
            sheet_rows = None
            try:
                from python_calamine import CalamineWorkbook
                wb = CalamineWorkbook.from_path(file_path)
                sheet_rows = list(wb.get_sheet_by_index(0).to_python())
            except Exception:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
                ws = wb.active
                sheet_rows = list(ws.iter_rows(values_only=True))
                wb.close()

            if not sheet_rows or len(sheet_rows) < 2:
                _rhi_set_job(job_id, status="error", message="Uploaded Excel file is empty")
                return

            headers = sheet_rows[0]
            idx = {str(h).strip().lower(): i for i, h in enumerate(headers)}
            for r in sheet_rows[1:]:
                if any(v is not None for v in r):
                    raw_rows.append([r[idx.get(c, -1)] if idx.get(c, -1) != -1 else "" for c in RHI_COLUMNS])

        if not raw_rows:
            _rhi_set_job(job_id, status="error", message="No valid data rows found in uploaded file")
            return

        _rhi_set_job(job_id, status="processing", pct=15, message=f"Sanitizing {len(raw_rows):,} rows...", stage="validate")

        converted = []
        weeknums = set()
        week_idx = RHI_COLUMNS.index("weeknum")

        for r in raw_rows:
            yw = _safe_int(r[week_idx], 0)
            if yw <= 0:
                continue
            weeknums.add(yw)
            out = []
            for i, val in enumerate(r):
                col_name = RHI_COLUMNS[i]
                if i == week_idx:
                    out.append(str(yw))
                elif col_name in RHI_INT_COLS:
                    out.append(str(_safe_int(val, 0)))
                elif col_name in RHI_FLOAT_COLS:
                    fv = _safe_float(val)
                    out.append("" if fv is None else str(fv))
                else:
                    out.append(_safe_str(val, ""))
            converted.append(out)

        if not converted:
            _rhi_set_job(job_id, status="error", message="No rows with valid weeknum found in file")
            return

        _rhi_set_job(job_id, status="processing", pct=25, message=f"Prepared {len(converted):,} rows. Uploading to database...", stage="copy")

        buf = io.StringIO()
        csv.writer(buf, lineterminator="\n").writerows(converted)
        payload = buf.getvalue().encode("utf-8")

        with get_postgres_connection() as conn:
            cur = conn.cursor()
            if weeknums:
                cur.execute(
                    "DELETE FROM sow.rhi_score WHERE weeknum IN (%s)" % ",".join(["%s"] * len(weeknums)),
                    tuple(weeknums)
                )
            col_names_sql = ",".join(['"%s"' % c for c in RHI_COLUMNS])
            cr = _RhiCountingReader(io.BytesIO(payload), len(payload), job_id)
            cur.copy_expert(
                f"COPY sow.rhi_score ({col_names_sql}) FROM STDIN WITH (FORMAT CSV, NULL '')",
                cr
            )
            conn.commit()
            cur.close()

        clear_rhi_caches()
        _rhi_set_job(job_id, status="success", pct=100, message=f"Successfully imported {len(converted):,} rows across week(s): {', '.join(map(str, sorted(weeknums)))}.")
    except Exception as e:
        logger.exception("RHI bulk import failed: %s", e)
        _rhi_set_job(job_id, status="error", pct=0, message=f"Import failed: {str(e)}")
    finally:
        if os.path.exists(file_path):
            try:
                os.unlink(file_path)
            except Exception:
                pass
